import csv
import io
import random
import re
import uuid
from datetime import datetime, timezone
from zipfile import BadZipFile

from core.db import get_store
from services import accounts, employee_tests, settings

PRODUCTS_KEY = "employee_portal_products"
MAX_QUESTION_BANK_SIZE = 5000


def product_key(product: str) -> str:
    return " ".join(str(product or "").split()).casefold()


def _products() -> dict:
    value = settings.get_setting(PRODUCTS_KEY)
    return value if isinstance(value, dict) else {}


def product_configs() -> list[dict]:
    configs = _products()
    out = []
    for key, config in sorted(configs.items(), key=lambda item: item[1].get("product", "").casefold()):
        version = config.get("bankVersion")
        question_count = get_store().count(
            "employee_portal_question_bank", {"product_key": key, "bank_version": version}
        ) if version else 0
        out.append({
            "product": config.get("product") or key,
            "questionCount": int(config.get("questionCount") or 0),
            "bankQuestionCount": question_count,
            "bankVersion": version,
        })
    return out


def _bank_questions(key: str, version: str) -> list[dict]:
    questions = get_store().select(
        "employee_portal_question_bank",
        {"product_key": key, "bank_version": version},
        ["question_text", "options", "correct_option_index", "explanation", "difficulty"],
    )
    if not questions:
        raise ValueError("The configured question bank is empty or unavailable.")
    return questions


def _cell(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return "" if value is None else str(value).strip()


def _header_key(value) -> str:
    return re.sub(r"[^a-z0-9]+", " ", _cell(value).casefold()).strip()


def _parse_wide_rows(rows: list[tuple], header: list[str]) -> list[dict]:
    def find(*names):
        return next((i for i, value in enumerate(header) if any(name in value for name in names)), None)

    q_idx = find("question")
    option_indices = [find(f"option {letter}", f"option{letter}") for letter in "abcd"]
    answer_idx = find("correct answer", "answer", "correct option", "correct")
    if q_idx is None or any(index is None for index in option_indices) or answer_idx is None:
        # Sample TPM books place the question in column B and options in C:F.
        q_idx, option_indices, answer_idx = 1, [2, 3, 4, 5], 6
    explanation_idx = find("explanation")
    letter_index = {letter.upper(): index for index, letter in enumerate("abcd")}
    output = []
    for row in rows:
        question = _cell(row[q_idx]) if q_idx < len(row) else ""
        options = [_cell(row[i]) if i < len(row) else "" for i in option_indices]
        answer = _cell(row[answer_idx]).casefold() if answer_idx < len(row) else ""
        if not question:
            continue
        answer_letter = answer[:1].upper()
        if answer_letter in letter_index:
            correct = letter_index[answer_letter]
        else:
            correct = next((i for i, option in enumerate(options) if option and option.casefold() == answer), None)
        if correct is None or any(not option for option in options):
            continue
        explanation = _cell(row[explanation_idx]) if explanation_idx is not None and explanation_idx < len(row) else ""
        output.append({
            "question_text": question, "options": options, "correct_option_index": correct,
            "explanation": explanation, "difficulty": "medium",
        })
    return output


def _parse_stacked_rows(rows: list[tuple], header: list[str]) -> list[dict]:
    q_idx = next((i for i, value in enumerate(header) if "question" in value), 3)
    option_idx = next((i for i, value in enumerate(header) if value.startswith("option")), 4)
    answer_idx = next((i for i, value in enumerate(header) if "answer" in value or "correct" in value or "✓" in value), 5)
    groups: dict[str, dict] = {}
    current_question = ""
    for row in rows:
        question = _cell(row[q_idx]) if q_idx < len(row) else ""
        option = _cell(row[option_idx]) if option_idx < len(row) else ""
        answer = _cell(row[answer_idx]).casefold() if answer_idx < len(row) else ""
        if question:
            current_question = question
        if not current_question or not option:
            continue
        group = groups.setdefault(current_question, {"options": [], "correct": None})
        if option not in group["options"]:
            group["options"].append(option)
        if answer.startswith("correct") or answer in ("✓", "✔", "yes", "true"):
            group["correct"] = group["options"].index(option)
    questions = []
    for question, group in groups.items():
        if group["correct"] is None or group["correct"] >= 4 or not 2 <= len(group["options"]) <= 4:
            continue
        options = group["options"][:4]
        options.extend(["N/A"] * (4 - len(options)))
        questions.append({
            "question_text": question,
            "options": options,
            "correct_option_index": group["correct"],
            "explanation": "Imported from the product question bank.",
            "difficulty": "medium",
        })
    return questions


def parse_question_bank(content: bytes, filename: str, sheet_name: str = "") -> list[dict]:
    suffix = filename.rsplit(".", 1)[-1].casefold()
    if suffix == "csv":
        try:
            rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        except UnicodeDecodeError as err:
            raise ValueError("The question-bank CSV must use UTF-8 encoding.") from err
        sheets = {"CSV": rows}
        selected = "CSV"
    elif suffix in ("xlsx", "xlsm"):
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException) as err:
            raise ValueError("The question-bank workbook could not be opened.") from err
        if not workbook.sheetnames:
            raise ValueError("The question-bank workbook has no sheets.")
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet {sheet_name!r} was not found. Available sheets: {', '.join(workbook.sheetnames)}")
        selected = sheet_name or workbook.sheetnames[0]
        sheets = {selected: list(workbook[selected].iter_rows(values_only=True))}
        workbook.close()
    else:
        raise ValueError("Upload an .xlsx, .xlsm, or .csv question bank.")

    rows = sheets[selected]
    if len(rows) < 2:
        raise ValueError("The selected question-bank sheet has no question rows.")
    header = [_header_key(value) for value in rows[0]]
    if any("option a" in value or "optiona" in value for value in header):
        questions = _parse_wide_rows(rows[1:], header)
    else:
        questions = _parse_stacked_rows(rows[1:], header)
        if not questions:
            questions = _parse_wide_rows(rows[1:], header)
    if not questions:
        raise ValueError("No complete multiple-choice questions with a marked correct answer were found.")
    if len(questions) > MAX_QUESTION_BANK_SIZE:
        raise ValueError(f"Question banks are limited to {MAX_QUESTION_BANK_SIZE} questions.")
    return questions


def save_question_bank(product: str, question_count: int, questions: list[dict]) -> dict:
    product = " ".join(product.split())
    key = product_key(product)
    if not product or len(product) > 200:
        raise ValueError("Product name is required and must be 200 characters or fewer.")
    if question_count < 1 or question_count > len(questions):
        raise ValueError(f"Question count must be between 1 and the bank size ({len(questions)}).")
    version = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    store = get_store()
    store.insert("employee_portal_question_bank", [
        {
            "id": str(uuid.uuid4()), "bank_version": version, "product_key": key,
            "question_text": question["question_text"], "options": question["options"],
            "correct_option_index": question["correct_option_index"],
            "explanation": question.get("explanation") or "", "difficulty": question.get("difficulty") or "medium",
            "created_at": now,
        }
        for question in questions
    ])
    configs = _products()
    configs[key] = {"product": product, "questionCount": question_count, "bankVersion": version}
    settings.set_setting(PRODUCTS_KEY, configs)
    return {"product": product, "questionCount": question_count, "bankQuestionCount": len(questions), "bankVersion": version}


def assign_product_test(employee: dict, config: dict, rng=None) -> dict:
    key = product_key(employee.get("product"))
    questions = _bank_questions(key, config["bankVersion"])
    count = int(config["questionCount"])
    if count > len(questions):
        raise ValueError(f"The question bank for {config['product']} has fewer than {count} questions.")
    chosen = (rng or random.SystemRandom()).sample(questions, count)
    store = get_store()
    employee_uuid = accounts.employee_uuid(employee["employee_id"])
    if not employee_uuid:
        raise RuntimeError(f"Could not resolve employee account {employee['employee_id']}.")
    test = store.first("tests", {"employee_id": employee_uuid, "topic_id": accounts.PRODUCT_ASSESSMENT_TOPIC_ID})
    now = datetime.now(timezone.utc).isoformat()
    time_limit = max(1800, count * 72)
    if test:
        employee_tests.archive_as_history(test["id"])
        store.delete("test_attempts", {"test_id": test["id"]})
        store.delete("test_questions", {"test_id": test["id"]})
        store.update("tests", {
            "employee_code": employee["employee_id"], "topic_title": config["product"],
            "subject_id": "resource-subject", "subject_title": "Product Assessment",
            "difficulty": "medium", "total_questions": count, "time_limit_seconds": time_limit,
            "status": "pending", "current_question_index": 0, "started_at": None, "completed_at": None,
            "in_progress": None, "session_recording_url": None, "proctoring": None,
            "score_correct": None, "score_total": None, "score_percent": None, "ai_analysis": None,
        }, {"id": test["id"]})
        test_id = test["id"]
    else:
        test_id = str(uuid.uuid4())
        store.insert("tests", {
            "id": test_id, "employee_id": employee_uuid, "employee_code": employee["employee_id"],
            "topic_id": accounts.PRODUCT_ASSESSMENT_TOPIC_ID, "subject_id": "resource-subject",
            "topic_title": config["product"], "subject_title": "Product Assessment", "difficulty": "medium",
            "total_questions": count, "time_limit_seconds": time_limit, "status": "pending",
            "current_question_index": 0, "created_at": now,
        })
    store.insert("test_questions", [
        {
            "id": str(uuid.uuid4()), "test_id": test_id, "question_index": index,
            "question_text": question["question_text"], "options": question["options"],
            "correct_option_index": question["correct_option_index"], "explanation": question.get("explanation") or "",
            "difficulty": question.get("difficulty") or "medium",
            "topic_id": accounts.PRODUCT_ASSESSMENT_TOPIC_ID, "topic_title": config["product"], "created_at": now,
        }
        for index, question in enumerate(chosen)
    ])
    return {"testId": test_id, "questionCount": count}


def _assignment_uses_bank(test_id: str, bank_questions: list[dict]) -> bool:
    assigned = get_store().select(
        "test_questions",
        {"test_id": test_id},
        ["question_text", "options", "correct_option_index"],
    )
    bank_items = {
        (
            question["question_text"],
            tuple(question["options"]),
            int(question["correct_option_index"]),
        )
        for question in bank_questions
    }
    return bool(assigned) and all(
        (
            question["question_text"],
            tuple(question["options"]),
            int(question["correct_option_index"]),
        ) in bank_items
        for question in assigned
    )


def provision_employee(details: dict) -> dict:
    employee_id = str(details.get("employee_id") or "").strip()
    full_name = str(details.get("full_name") or "").strip()
    email = str(details.get("email") or "").strip().lower()
    product = " ".join(str(details.get("product") or "").split())
    designation = str(details.get("designation") or "").strip()
    department = str(details.get("department") or "").strip()
    if not employee_id or len(employee_id) > 40 or not re.fullmatch(r"[A-Za-z0-9._-]+", employee_id):
        raise ValueError("Employee ID is required and may contain only letters, numbers, periods, underscores, and hyphens.")
    if not full_name or len(full_name) > 120:
        raise ValueError("Name is required and must be 120 characters or fewer.")
    if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email) or len(email) > 254:
        raise ValueError("A valid employee email address is required.")
    if not product:
        raise ValueError("Product is required.")
    configs = _products()
    config = configs.get(product_key(product))
    if not config:
        raise ValueError(f"No question bank is configured for product {product!r}.")
    existing = accounts.get_account(employee_id)
    if existing and existing.get("role") != "employee":
        raise ValueError("This Employee ID is already assigned to a non-employee account.")
    existing_test = None
    bank_questions = None
    if existing:
        employee_uuid = accounts.employee_uuid(employee_id)
        if employee_uuid:
            existing_test = get_store().first(
                "tests", {"employee_id": employee_uuid, "topic_id": accounts.PRODUCT_ASSESSMENT_TOPIC_ID},
                ["id", "status", "topic_title", "total_questions"],
            )
            if existing_test:
                bank_questions = _bank_questions(product_key(config["product"]), config["bankVersion"])
    product_changed = bool(existing and product_key(existing.get("product")) != product_key(config["product"]))
    assignment_matches = bool(
        existing_test
        and not product_changed
        and product_key(existing_test.get("topic_title")) == product_key(config["product"])
        and int(existing_test.get("total_questions") or 0) == int(config["questionCount"])
        and get_store().count("test_questions", {"test_id": existing_test["id"]}) == int(config["questionCount"])
        and bank_questions is not None
        and _assignment_uses_bank(existing_test["id"], bank_questions)
    )
    if existing_test and not assignment_matches and existing_test.get("status") == "in_progress":
        raise ValueError("This employee has an in-progress assessment. Finish it before changing the product assignment.")
    if existing:
        needs_initial_password = not accounts.has_password(existing)
        default_hash, default_salt = accounts.hash_password(employee_id) if needs_initial_password else (None, None)
        default_password_active = needs_initial_password or accounts.verify_password(
            employee_id, existing.get("password_salt") or "", existing.get("password_hash") or ""
        )
        account = {
            **existing, "full_name": full_name, "email": email, "designation": designation,
            "department": department or existing.get("department") or "general",
            "product": config["product"], "role": "employee", "product_qb_eligible": True, "assessment_only": True,
        }
        if existing.get("email", "").strip().casefold() != email.casefold():
            account["email_verified"] = False
        if needs_initial_password:
            account.update({
                "password_hash": default_hash, "password_salt": default_salt, "is_first_login": False,
            })
        account["is_first_login"] = False
    else:
        needs_initial_password = True
        default_password_active = True
        password_hash, password_salt = accounts.hash_password(employee_id)
        account = {
            "employee_id": employee_id, "full_name": full_name, "email": email, "designation": designation,
            "department": department or "general", "role": "employee", "product": config["product"],
            "product_qb_eligible": True, "assessment_only": True, "is_first_login": False, "email_verified": False,
            "password_hash": password_hash, "password_salt": password_salt,
            "xp_points": 0, "streak_days": 0, "skill_level": "beginner", "ai_readiness_score": 0,
        }
    if not accounts.upsert_account(account):
        raise RuntimeError(f"Could not save employee account {employee_id}.")
    saved = accounts.get_account(employee_id)
    if not saved:
        raise RuntimeError(f"Employee account {employee_id} was not available after saving.")
    if assignment_matches and existing_test:
        assignment = {"testId": existing_test["id"], "questionCount": int(existing_test["total_questions"])}
    else:
        try:
            assignment = assign_product_test(saved, config)
        except Exception:
            if not existing:
                get_store().delete("employees", {"employee_id": employee_id})
            raise
    return {
        "employeeId": saved["employee_id"], "name": saved["full_name"], "email": saved["email"],
        "designation": saved["designation"], "product": saved["product"],
        "firstLoginPassword": employee_id if default_password_active else None,
        "testId": assignment["testId"], "questionCount": assignment["questionCount"],
        "created": existing is None,
    }


def parse_employee_upload(content: bytes, filename: str) -> list[dict]:
    suffix = filename.rsplit(".", 1)[-1].casefold()
    if suffix == "csv":
        try:
            rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        except UnicodeDecodeError as err:
            raise ValueError("The roster CSV must use UTF-8 encoding.") from err
    elif suffix in ("xlsx", "xlsm"):
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException) as err:
            raise ValueError("The employee roster workbook could not be opened.") from err
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
    else:
        raise ValueError("Upload an .xlsx, .xlsm, or .csv employee roster.")
    if not rows:
        raise ValueError("The employee roster is empty.")
    headers = [_header_key(cell) for cell in rows[0]]

    def column(*names):
        return next((index for index, header in enumerate(headers) if header in names), None)

    indices = {
        "employee_id": column("emp id", "employee id", "tpm id employee id", "id"),
        "full_name": column("name", "emp name", "full name"),
        "email": column("nokia email id", "email", "email id", "employee email"),
        "designation": column("role", "designation", "job role"),
        "product": column("product"),
    }
    if any(indices[name] is None for name in ("employee_id", "full_name", "email", "product")):
        raise ValueError("Roster columns must include Emp ID, Name, Email, and Product. Role/Designation is optional.")
    records = []
    for line, row in enumerate(rows[1:], start=2):
        if not any(_cell(value) for value in row):
            continue
        records.append({
            **{name: _cell(row[index]) if index is not None and index < len(row) else "" for name, index in indices.items()},
            "_row": line,
        })
    if not records:
        raise ValueError("The employee roster has no employee rows.")
    if len(records) > 500:
        raise ValueError("Upload no more than 500 employees at a time.")
    return records
