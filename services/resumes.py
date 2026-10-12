import base64
import hashlib
import io
import json
import logging
import re
import threading
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from typing import Any, Optional

from core.db import get_store
from services import ai, sessions, storage

log = logging.getLogger("talentscope")

_cache: dict[str, dict] = {}
_cache_lock = threading.Lock()
_executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="resume-ai")

TECH_SKILLS = [
    "javascript", "typescript", "python", "react", "vue", "angular", "node", "express", "java", "spring", "c#", ".net",
    "go", "rust", "php", "docker", "kubernetes", "aws", "azure", "gcp", "postgresql", "mongodb", "mysql", "redis", "git",
    "ci/cd", "html", "css", "sass", "webpack", "linux", "bash",
]

HEADERS = [
    ("summary", r"^(?:profile\s+summary|summary|profile|professional\s+summary|objective|career\s+objective|about\s+me|about)$"),
    ("experience", r"^(?:experience|professional\s+experience|work\s+experience|employment|employment\s+history|work\s+history|internships?|internship\s+experience|professional\s+history)$"),
    ("education", r"^(?:education|academic\s+background|academic\s+profile|academics?|education\s+details?|academic\s+details?)$"),
    ("skills", r"^(?:skills|technical\s+skills|core\s+competencies|key\s+skills|expertise|skills\s+&\s+tools|technologies)$"),
    ("projects", r"^(?:projects|academic\s+projects|personal\s+projects|selected\s+projects)$"),
    ("certifications", r"^(?:certifications?|licenses\s+&\s+certifications?|credentials?)$"),
    ("publications", r"^(?:publications?|research\s+publications?|papers)$"),
]

MONTH = r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*"
DATE_RANGE = re.compile(
    rf"(?:{MONTH}\s+)?\b(?:19|20)\d{{2}}\s*(?:-|–|—|to)\s*(?:(?:(?:{MONTH}\s+)?\b(?:19|20)\d{{2}})|Present|Current)\b", re.I
)
TRIM_RE = re.compile(r"^[•*\-\s,]+|[•*\-\s,]+$")


def _uid() -> str:
    return str(uuid.uuid4())


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _normalize_text(text: str) -> str:
    out = (text or "").replace("\x00", "")
    out = re.sub(r"[ \t]*\n[ \t]*", "\n", out)
    out = re.sub(r"[ \t]{2,}", " ", out)
    return re.sub(r"\n{3,}", "\n\n", out).strip()


def _strip_html(html: str) -> str:
    out = re.sub(r"<script[\s\S]*?</script>", " ", html, flags=re.I)
    out = re.sub(r"<style[\s\S]*?</style>", " ", out, flags=re.I)
    out = re.sub(r"<br\s*/?>", "\n", out, flags=re.I)
    out = re.sub(r"</(p|div|h\d|li|tr)>", "\n", out, flags=re.I)
    out = re.sub(r"<[^>]+>", " ", out)
    for src, dst in (("&nbsp;", " "), ("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">")):
        out = re.sub(src, dst, out, flags=re.I)
    out = re.sub(r"[ \t]+\n", "\n", out)
    out = re.sub(r"\n{3,}", "\n\n", out)
    return re.sub(r"[ \t]{2,}", " ", out).strip()


def _extract_pdf(buffer: bytes) -> str:
    from pypdf import PdfReader

    text = ""
    try:
        reader = PdfReader(io.BytesIO(buffer))
        text = _normalize_text("\n".join((page.extract_text() or "") for page in reader.pages))
    except Exception as err:
        log.warning("pypdf extraction failed: %s", err)
    if text.strip():
        return text
    raise ValueError("Could not extract text from this PDF. Try a text-based PDF or a DOCX.")


def _extract_docx(buffer: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(buffer))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as err:
        raise ValueError(f"DOCX parsing failed: {err}")


def extract_text(buffer: bytes, filename: Optional[str] = None) -> str:
    if not buffer:
        return ""
    ext = (filename or "").rsplit(".", 1)[-1].lower() if filename and "." in filename else ""
    head = buffer[:1024].decode("utf-8", errors="ignore").lstrip("\ufeff").lower()
    is_pdf = buffer[:2] == b"%P"
    is_zip = buffer[:2] == b"PK"
    looks_html = (
        ext in ("html", "htm") or "<html" in head or "<!doctype html" in head or "<head" in head or "<body" in head
        or re.search(r"<(div|p|table|meta|span|h[1-6])[\s>]", head) is not None
    )
    if looks_html and not is_pdf and not is_zip:
        return _strip_html(buffer.decode("utf-8", errors="ignore").lstrip("\ufeff"))
    if is_pdf:
        return _extract_pdf(buffer)
    if is_zip:
        return _extract_docx(buffer)
    return buffer.decode("utf-8", errors="replace")


def _match_header(line: str) -> Optional[str]:
    clean = re.sub(r"[•*\-\u2022\s]+", " ", line).strip().lower()
    if not clean or len(clean) > 40:
        return None
    for name, pattern in HEADERS:
        if re.match(pattern, clean):
            return name
    return None


def _personal(lines: list[str]) -> dict:
    top = "\n".join(lines[:20])
    email = re.search(r"[\w.-]+@[\w.-]+\.\w+", top)
    phone = re.search(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3,5}\)?[-.\s]?\d{3,5}[-.\s]?\d{4}", top)
    linkedin = re.search(r"linkedin\.com/in/[\w-]+", top, re.I)
    github = re.search(r"github\.com/[\w-]+", top, re.I)
    return {
        "fullName": re.sub(r"[•*\-]", "", lines[0]).strip() if lines else "",
        "email": email.group(0) if email else "",
        "phone": phone.group(0) if phone else "",
        "linkedin": f"https://{linkedin.group(0)}" if linkedin else "",
        "github": f"https://{github.group(0)}" if github else "",
        "title": lines[1].strip() if len(lines) > 1 else "",
    }


def _is_bullet(line: str) -> bool:
    return line.startswith(("•", "-", "*"))


def _parse_experience(content: str) -> list[dict]:
    if not content:
        return []
    items: list[dict] = []
    current: Optional[dict] = None
    for line in (ln.strip() for ln in content.split("\n")):
        if not line:
            continue
        match = DATE_RANGE.search(line)
        if match:
            if current:
                items.append(current)
            parts = re.split(r"\s*(?:-|–|—|to)\s*", match.group(0), flags=re.I)
            start = parts[0].strip() if parts else ""
            end = parts[1].strip() if len(parts) > 1 else ""
            clean_line = DATE_RANGE.sub("", line, count=1)
            clean_line = re.sub(rf"\s*{MONTH}\s*$", "", clean_line, flags=re.I).strip()
            pieces = re.split(r"\s*(?:—|–|\s-\s|@|\bat\b)\s*", clean_line)
            position, company = clean_line, ""
            if len(pieces) >= 2:
                position = pieces[0].strip()
                company = pieces[1].strip()
                if len(pieces) > 2:
                    company += ", " + ", ".join(pieces[2:])
            current = {
                "id": _uid(), "company": TRIM_RE.sub("", company).strip(), "position": TRIM_RE.sub("", position).strip(),
                "startDate": start, "endDate": end, "current": bool(re.search(r"Present|Current", end, re.I)),
                "description": "", "bulletPoints": [], "technologies": [],
            }
            continue
        if _is_bullet(line):
            if current is not None:
                current["bulletPoints"].append(
                    {"id": _uid(), "text": re.sub(r"^[•*\-]\s*", "", line), "impact": "medium", "issues": []}
                )
            continue
        if current is not None:
            current["description"] += f"{line} "
    if current:
        items.append(current)
    return items


def _parse_education(content: str) -> list[dict]:
    if not content:
        return []
    out = []
    for line in (ln for ln in content.split("\n") if ln):
        year = re.search(r"\b(?:19|20)\d{2}\b", line)
        out.append({
            "id": _uid(), "institution": line.split(",")[0].strip(), "degree": line.strip(),
            "graduationDate": year.group(0) if year else "",
        })
    return out


def _parse_projects(content: str) -> list[dict]:
    if not content:
        return []
    projects: list[dict] = []
    current: Optional[dict] = None
    for line in (ln.strip() for ln in content.split("\n")):
        if not line:
            continue
        if not _is_bullet(line):
            if current:
                projects.append(current)
            name = re.split(r"\s*(?:—|–|\s-\s)\s*", line)[0].strip()
            current = {"id": _uid(), "name": name, "description": line, "technologies": [], "bulletPoints": []}
            continue
        if current is not None:
            current["bulletPoints"].append(re.sub(r"^[•*\-]\s*", "", line))
    if current:
        projects.append(current)
    return projects


def _skills(text: str) -> dict:
    lower = text.lower()
    return {
        "technical": [s for s in TECH_SKILLS if s in lower],
        "soft": ["communication", "leadership", "teamwork", "problem-solving"],
        "tools": ["git", "docker", "jira"],
        "languages": [],
        "other": [],
    }


def parse_structure(text: str) -> dict:
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    boundaries: list[dict] = []
    for idx, line in enumerate(lines):
        name = _match_header(line)
        if name and (not boundaries or boundaries[-1]["lineIndex"] != idx):
            boundaries.append({"name": name, "lineIndex": idx})
    sections: dict[str, str] = {}
    section_list = []
    for i, b in enumerate(boundaries):
        end = boundaries[i + 1]["lineIndex"] if i + 1 < len(boundaries) else len(lines)
        content = "\n".join(lines[b["lineIndex"] + 1:end]).strip()
        sections[b["name"]] = (sections[b["name"]] + "\n" + content) if sections.get(b["name"]) else content
        section_list.append({"name": b["name"], "confidence": 0.9, "startIndex": b["lineIndex"], "endIndex": end, "content": content})
    certs = []
    if sections.get("certifications"):
        for line in (ln for ln in sections["certifications"].split("\n") if ln):
            certs.append({"id": _uid(), "name": re.sub(r"[•*\-]", "", line).strip(), "issuer": "Self", "date": ""})
    return {
        "personal": _personal(lines),
        "summary": sections.get("summary", ""),
        "experience": _parse_experience(sections.get("experience", "")),
        "education": _parse_education(sections.get("education", "")),
        "skills": _skills(text),
        "projects": _parse_projects(sections.get("projects", "")),
        "certifications": certs,
        "achievements": [],
        "leadership": [],
        "sections": section_list,
    }


ROLE_DEFS = [
    ("Frontend Developer", ["react", "vue", "angular", "html", "css", "javascript", "typescript", "frontend", "ui", "web"]),
    ("Backend Developer", ["node", "python", "java", "spring", "c#", ".net", "go", "ruby", "django", "flask", "backend", "api", "database", "sql", "postgresql", "mysql", "mongodb"]),
    ("Full Stack Developer", ["react", "node", "python", "java", "vue", "angular", "full stack", "fullstack", "javascript", "typescript"]),
    ("Machine Learning Engineer", ["machine learning", "ml", "pytorch", "tensorflow", "model", "deep learning", "nlp", "computer vision"]),
    ("AI Engineer", ["ai", "artificial intelligence", "openai", "llm", "langchain", "prompt", "transformer", "hugging face", "agent"]),
    ("RAG Specialist", ["rag", "pinecone", "vector", "qdrant", "chroma", "retrieval", "embedding"]),
    ("Data Scientist", ["python", "r", "machine learning", "ml", "ai", "pandas", "numpy", "tensorflow", "pytorch", "data science", "statistics", "data analysis"]),
    ("Data Engineer", ["python", "sql", "etl", "spark", "hadoop", "kafka", "data pipeline", "airflow", "aws", "gcp"]),
    ("DevOps Engineer", ["docker", "kubernetes", "aws", "azure", "gcp", "ci/cd", "jenkins", "terraform", "ansible", "linux", "bash", "devops"]),
    ("UI/UX Designer", ["figma", "sketch", "adobe", "ui", "ux", "wireframe", "prototype", "user experience", "user interface", "design"]),
    ("Product Manager", ["product", "roadmap", "agile", "scrum", "jira", "stakeholder", "strategy", "product management", "market research"]),
]


def _roles(parsed: dict) -> list[str]:
    sk = parsed.get("skills") or {}
    skills = [str(s).lower() for key in ("technical", "soft", "tools", "languages", "other") for s in (sk.get(key) or [])]
    title = str((parsed.get("personal") or {}).get("title") or "").lower()
    exp_titles = " ".join(str(e.get("position") or "").lower() for e in parsed.get("experience") or [])
    scored = []
    for role, keywords in ROLE_DEFS:
        score = matched = 0
        for kw in keywords:
            if any(kw in s for s in skills):
                score += 10
                matched += 1
            if kw in title:
                score += 25
            if kw in exp_titles:
                score += 10
        if role == "Full Stack Developer":
            front = any(s in ("react", "vue", "angular", "html", "css") for s in skills)
            back = any(s in ("node", "python", "java", "sql", "database", "api") for s in skills)
            if not front or not back:
                score = 0
        scored.append((role, score, matched))
    top = sorted([r for r in scored if r[1] > 0 and r[2] >= 1], key=lambda r: -r[1])
    years = len(parsed.get("experience") or []) * 1.5
    result = []
    for role, _, _ in top[:3]:
        result.append(f"Senior {role}" if years >= 5 and "Manager" not in role and "Specialist" not in role else role)
    if any("AI" in r for r in result) and any("Full Stack" in r for r in result):
        result.insert(0, "Full-Stack AI Developer")
    return list(dict.fromkeys(result))[:4] if result else ["Professional"]


def _industry_fit(parsed: dict) -> list[dict]:
    tech = [t.lower() for t in (parsed.get("skills") or {}).get("technical", [])]
    defs = [
        ("Software Development", ["javascript", "python", "react", "node"], "Strong software development stack"),
        ("Cloud / DevOps", ["aws", "docker", "kubernetes", "ci/cd"], "Strong infrastructure skills"),
        ("Data Engineering", ["python", "sql", "etl", "analytics"], "Strong data engineering skillset"),
    ]
    fits = [{"industry": n, "matchScore": sum(1 for t in tech if t in keys) * 25, "rationale": r} for n, keys, r in defs]
    return sorted(fits, key=lambda f: -f["matchScore"])[:3]


def _executive(score: float, parsed: dict) -> str:
    top = ", ".join(((parsed.get("skills") or {}).get("technical") or [])[:3])
    if score >= 80:
        return (f"A highly proficient professional with a strong background in {top}. The profile demonstrates clear quantifiable impact and exceptional capabilities."
                if top else "Strong resume with well-structured experience and quantifiable achievements.")
    if score >= 60:
        return (f"A capable professional skilled in {top}. The resume has a good foundation but could benefit from stronger action verbs and clearer metrics."
                if top else "Good resume foundation but improvements to quantifiable impact and wording are recommended.")
    return "Resume requires significant optimization, specifically around measurable outcomes and formatting."


def _insights(analysis: dict) -> list[str]:
    scores = analysis.get("scores") or {}
    out = [f"Resume Score: {analysis.get('overallScore') or 0}/100"]
    if (scores.get("actionVerbs") if scores.get("actionVerbs") is not None else 100) < 60:
        out.append("Use stronger action verbs.")
    if (scores.get("measurability") if scores.get("measurability") is not None else 100) < 50:
        out.append("Add measurable achievements.")
    if (analysis.get("atsScore") if analysis.get("atsScore") is not None else 100) < 70:
        out.append("Improve ATS formatting.")
    return out


def _num(value: Any, default: float) -> float:
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) and value else default


def generate_report(parsed: dict, analysis: dict) -> dict:
    score = _num(analysis.get("overallScore"), 50)
    confidence = analysis.get("hiringConfidence") or (
        "very-high" if score >= 85 else "high" if score >= 70 else "medium" if score >= 50 else "low")
    scores = analysis.get("scores") or {}
    suggestions = ai.generate_suggestions(analysis)[:5]
    jd = analysis.get("jdMatchScore")
    return {
        "executiveSummary": analysis.get("executiveSummary") or _executive(score, parsed),
        "recruiterInsights": _insights(analysis),
        "hiringConfidence": confidence,
        "industryFit": analysis.get("industryFit") or _industry_fit(parsed),
        "targetRoles": analysis.get("targetRoles") or _roles(parsed),
        "priorityImprovements": [
            {"rank": i + 1, "category": s["section"], "title": str(s["description"]).split(".")[0],
             "description": s["description"], "effort": "medium", "impact": "high" if s["priority"] == "high" else "medium"}
            for i, s in enumerate(suggestions)
        ],
        "visualMetrics": {
            "overallScore": score,
            "atsScore": _num(analysis.get("atsScore"), 50),
            "radarData": [
                {"subject": "Content", "value": _num(scores.get("clarity"), 50), "fullMark": 100},
                {"subject": "Format", "value": _num(scores.get("formatting"), 50), "fullMark": 100},
                {"subject": "ATS", "value": _num(analysis.get("atsScore"), 50), "fullMark": 100},
                {"subject": "Impact", "value": _num(analysis.get("impactScore"), 50), "fullMark": 100},
                {"subject": "Technical", "value": _num(analysis.get("technicalScore"), 50), "fullMark": 100},
            ],
            "scoreHistory": [],
            "topSkills": [{"name": n, "score": max(50, 100 - i * 10)} for i, n in enumerate(((parsed.get("skills") or {}).get("technical") or [])[:5])],
        },
        "suitability": analysis.get("suitability") or ("suitable" if jd is not None and jd >= 40 else "unsuitable"),
        "jdMatchScore": jd if jd is not None else None,
        "jdMatchRationale": analysis.get("jdMatchRationale") or None,
    }


def build_enhanced(parsed: dict, analysis: dict) -> dict:
    return {
        "summary": parsed.get("summary") or "",
        "experience": {
            e["id"]: {"bulletPoints": [{"original": b["text"], "enhanced": b["text"], "changes": []} for b in e.get("bulletPoints", [])]}
            for e in parsed.get("experience", [])
        },
        "projects": {
            p["id"]: {"description": p["description"], "bulletPoints": [{"original": b, "enhanced": b} for b in p.get("bulletPoints", [])]}
            for p in parsed.get("projects", [])
        },
        "skills": {"added": (analysis.get("keywordAnalysis") or {}).get("suggestedKeywords") or [], "removed": [], "reorganized": True},
        "suggestions": ai.generate_suggestions(analysis),
    }


def _json_field(value: Any) -> Any:
    if value is None:
        return {}
    if isinstance(value, (dict, list)):
        return value
    if isinstance(value, str):
        try:
            return json.loads(value)
        except ValueError:
            return {}
    return {}


def _map_row(row: dict) -> dict:
    created = row.get("created_at") or _now()
    return {
        "id": row.get("id"),
        "filename": row.get("filename") or "unknown",
        "originalText": row.get("text_content") or "",
        "parsed": _json_field(row.get("parsed")),
        "analysis": _json_field(row.get("analysis")),
        "enhanced": _json_field(row.get("enhanced")),
        "report": _json_field(row.get("report")),
        "error": row.get("error") or None,
        "fileHash": row.get("file_hash") or None,
        "fileBase64": row.get("file_base64") or None,
        "createdAt": created,
        "updatedAt": row.get("updated_at") or created,
        "status": "failed" if row.get("error") else "completed",
    }


def save_row(resume: dict) -> None:
    def dump(value):
        return json.dumps(value) if value else None

    row = {
        "id": resume["id"],
        "filename": resume.get("filename") or "unknown",
        "text_content": resume.get("originalText") or "",
        "parsed": dump(resume.get("parsed")),
        "analysis": dump(resume.get("analysis")),
        "enhanced": dump(resume.get("enhanced")),
        "report": dump(resume.get("report")),
        "error": resume.get("error"),
        "file_hash": resume.get("fileHash"),
        "file_base64": resume.get("fileBase64"),
    }
    with _cache_lock:
        _cache[resume["id"]] = resume
    try:
        get_store().upsert("resumes", row, "id")
    except Exception as err:
        log.error("Failed to save resume row: %s", err)


def get_resume(resume_id: str, force_fresh: bool = False) -> Optional[dict]:
    if not force_fresh:
        with _cache_lock:
            cached = _cache.get(resume_id)
        if cached:
            return cached
    row = get_store().first("resumes", {"id": resume_id})
    if not row:
        return None
    resume = _map_row(row)
    with _cache_lock:
        _cache[resume_id] = resume
    return resume


def list_resumes() -> list[dict]:
    rows: list[dict] = []
    page, offset = 500, 0
    cols = ["id", "filename", "parsed", "report", "error", "created_at", "file_hash"]
    store = get_store()
    while True:
        batch = store.select("resumes", columns=cols, order="created_at desc", limit=page, offset=offset)
        rows.extend(batch)
        if len(batch) < page:
            break
        offset += page
    out = []
    for row in rows:
        try:
            out.append(_map_row(row))
        except Exception as err:
            log.warning("Skipping corrupt resume row %s: %s", row.get("id"), err)
    return out


def _find_by_hash(file_hash: str) -> Optional[dict]:
    row = get_store().first("resumes", {"file_hash": file_hash}, ["id"], "created_at desc")
    return get_resume(row["id"]) if row else None


def find_by_filename(filename: str, rm_email: str) -> Optional[dict]:
    for row in get_store().select("resumes", {"filename": filename}, order="created_at desc"):
        report = _json_field(row.get("report"))
        if str(report.get("rmEmail") or "").lower().strip() == rm_email.lower().strip():
            return _map_row(row)
    return None


def delete_resume(resume_id: str) -> None:
    store = get_store()
    with _cache_lock:
        _cache.pop(resume_id, None)
    storage.remove_prefix("resumes", f"{resume_id}-")
    for table in ("interview_questions", "interview_attempts"):
        try:
            store.delete(table, {"resume_id": resume_id})
        except Exception as err:
            log.warning("Failed clearing %s for %s: %s", table, resume_id, err)
    store.delete("resumes", {"id": resume_id})
    try:
        sessions.delete_session_by_resume_id(resume_id)
    except Exception as err:
        log.warning("Failed to delete candidate session during resume deletion: %s", err)


def clear_all() -> None:
    with _cache_lock:
        _cache.clear()
    storage.clear_bucket("resumes")
    store = get_store()
    for table in ("interview_questions", "interview_attempts"):
        try:
            store.delete(table, None)
        except Exception as err:
            log.warning("Failed clearing %s: %s", table, err)
    store.delete("resumes", None)


def _complete_processing(resume: dict, buffer: bytes, jd_text: Optional[str], jd_id: Optional[str], rm_email: Optional[str]) -> None:
    try:
        text = extract_text(buffer, resume.get("filename"))
        resume["originalText"] = text
        parsed = parse_structure(text)
        resume["parsed"] = parsed
        analysis = ai.analyze_resume(text, parsed, jd_text)
        resume["analysis"] = analysis
        resume["enhanced"] = build_enhanced(parsed, analysis)
        report = generate_report(parsed, analysis)
        if jd_id:
            report["jdId"] = jd_id
        if rm_email:
            report["rmEmail"] = rm_email
        resume["report"] = report
        resume["status"] = "completed"
        resume["error"] = None
    except Exception as err:
        log.error("Resume processing failed: %s", err)
        resume["status"] = "failed"
        resume["error"] = str(err) or "Processing failed."
    resume["updatedAt"] = _now()
    save_row(resume)


def _update_existing(existing: dict, jd_id: Optional[str], rm_email: Optional[str]) -> dict:
    if (jd_id or rm_email) and existing.get("report") is not None:
        report = existing["report"]
        updated = False
        if jd_id and report.get("jdId") != jd_id:
            report["jdId"] = jd_id
            updated = True
        if rm_email and report.get("rmEmail") != rm_email:
            report["rmEmail"] = rm_email
            updated = True
        if updated:
            save_row(existing)
    return existing


def process_resume(filename: str, buffer: bytes, jd_text: Optional[str] = None, jd_id: Optional[str] = None,
                   rm_email: Optional[str] = None, force_replace: bool = False, wait: bool = False) -> dict:
    file_hash = hashlib.sha256(buffer).hexdigest()
    existing = _find_by_hash(file_hash)
    if not existing and filename and rm_email:
        existing = find_by_filename(filename, rm_email)
    if existing and not force_replace and (not wait or existing.get("status") == "completed"):
        return _update_existing(existing, jd_id, rm_email)
    resume_id = existing["id"] if existing and force_replace else _uid()
    if existing and force_replace:
        for table in ("interview_questions", "interview_attempts"):
            try:
                get_store().delete(table, {"resume_id": resume_id})
            except Exception as err:
                log.error("Failed to clear %s during forceReplace: %s", table, err)
    now = _now()
    resume = {
        "id": resume_id, "filename": filename, "originalText": "", "fileHash": file_hash,
        "fileBase64": base64.b64encode(buffer).decode("ascii"), "parsed": {}, "analysis": {}, "enhanced": {}, "report": {},
        "createdAt": now, "updatedAt": now, "status": "processing", "error": None,
    }
    resume["filePath"] = storage.save("resumes", f"{resume_id}-{filename}", buffer)
    save_row(resume)
    if wait:
        _complete_processing(resume, buffer, jd_text, jd_id, rm_email)
    else:
        _executor.submit(_complete_processing, resume, buffer, jd_text, jd_id, rm_email)
    return resume
