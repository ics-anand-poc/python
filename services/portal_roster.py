import io
import logging
import re
from typing import Optional

from core.db import get_store
from services import docs_storage
from services.products import format_product_display_name, format_topic_title_for_display

log = logging.getLogger("talentscope")
UUID_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)
PRODUCT_TOPIC = "resource-product-assessment"
HISTORY_PREFIX = "resource-product-assessment-history"
_mapping_cache: dict = {"key": None, "rows": []}


def normalize_employee_id(value) -> str:
    text = str(value if value is not None else "").strip()
    if re.fullmatch(r"\d+\.0+", text):
        text = text[: text.index(".")]
    return text


def is_employee_uuid(value) -> bool:
    return bool(UUID_RE.match(str(value or "").strip()))


def display_employee_code(code=None, owner=None, linked_code=None, linked_name=None) -> str:
    c, l, o = normalize_employee_id(code), normalize_employee_id(linked_code), normalize_employee_id(owner)
    if c and not is_employee_uuid(c):
        return c
    if l and not is_employee_uuid(l):
        return l
    if o and not is_employee_uuid(o):
        return o
    from_name = normalize_employee_id(linked_name)
    if re.fullmatch(r"\d{4,}", from_name):
        return from_name
    return l or c or o


def derive_status(assigned_count: int, test_id, raw_status, answered: int, total: int, started_at) -> str:
    if raw_status == "completed":
        return "completed"
    progress = round(answered / total * 100) if total > 0 else 0
    if not (assigned_count > 0 or test_id):
        return "unassigned"
    if 0 < progress < 100:
        return "pending"
    if raw_status == "in_progress" or started_at:
        return "in_progress"
    return "not_started"


def _cell(value) -> str:
    return "" if value is None else str(value).strip()


def _pick(get, keys: list[str]) -> str:
    for key in keys:
        value = get(key)
        if value:
            return value
    return ""


def parse_mapping_workbook(data: bytes) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.worksheets[0] if workbook.worksheets else None
    if sheet is None:
        return []
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []
    col = {_cell(h): i for i, h in enumerate(header) if _cell(h)}
    question_cols = [col[f"Assigned Question {n}"] for n in range(1, 26) if f"Assigned Question {n}" in col]
    by_id: dict[str, dict] = {}
    for row in rows:
        def get(key: str) -> str:
            idx = col.get(key)
            return _cell(row[idx]) if idx is not None and idx < len(row) else ""

        employee_id = _pick(get, ["Emp ID", "Employee ID"])
        if not employee_id:
            continue
        questions, seen = [], set()
        for idx in question_cols:
            question = _cell(row[idx]) if idx < len(row) else ""
            key = re.sub(r"\s+", " ", question.strip().lower())
            if question and key not in seen:
                seen.add(key)
                questions.append(question)
        entry = {
            "employee_id": employee_id, "full_name": _pick(get, ["Emp Name", "Employee Name"]), "role": get("Role"),
            "domain": get("Domain"), "product": _pick(get, ["Product", "Product-Updated"]), "email": _pick(get, ["Nokia Email ID", "Email"]),
            "ddh": _pick(get, ["DDH", "DDH Manager"]), "emp_status": get("Emp Status"), "remarks": get("Remarks"),
            "assigned_questions": questions, "assigned_question_count": len(questions),
        }
        key = normalize_employee_id(employee_id)
        existing = by_id.get(key)
        if not existing:
            by_id[key] = entry
            continue
        merged_q = list(dict.fromkeys([*existing["assigned_questions"], *entry["assigned_questions"]]))
        by_id[key] = {**existing, **{k: (v or existing[k]) for k, v in entry.items() if k not in ("assigned_questions", "assigned_question_count")},
                      "assigned_questions": merged_q, "assigned_question_count": len(merged_q)}
    return list(by_id.values())


def mapping_looks_valid(data: bytes) -> bool:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            for n, row in enumerate(sheet.iter_rows(min_row=1, max_row=5, values_only=True), 1):
                if any(re.search(r"assigned question\s*\d+", _cell(v), re.I) for v in row):
                    return True
    except Exception:
        return False
    return False


def invalidate_mapping_cache() -> None:
    _mapping_cache["key"] = None


def _mapping_file() -> Optional[str]:
    usable = [n for n in docs_storage.list_files("Portal Mapping") if re.search(r"\.(xlsx|xls)$", n, re.I) and not docs_storage.is_portal_credentials_name(n)]
    return next((n for n in usable if docs_storage.is_portal_mapping_name(n)), usable[0] if usable else None)


def load_mapping_rows() -> list[dict]:
    name = _mapping_file()
    if not name:
        return []
    key = (name, docs_storage.mtime("Portal Mapping", name))
    if _mapping_cache["key"] == key:
        return _mapping_cache["rows"]
    try:
        data = docs_storage.read_file("Portal Mapping", name)
        rows = parse_mapping_workbook(data) if data else []
    except Exception as err:
        log.warning("Portal mapping workbook unreadable: %s", err)
        rows = []
    _mapping_cache.update(key=key, rows=rows)
    return rows


def _questions_from_active_test(employee_id: str) -> list[str]:
    """Prefer live product-assessment questions over the mapping workbook snapshot."""
    store = get_store()
    emp_id = normalize_employee_id(employee_id)
    if not emp_id:
        return []
    tests = store.query(
        """
        SELECT TOP 1 t.id
        FROM dbo.tests t
        WHERE (t.employee_code = ? OR CONVERT(nvarchar(64), t.employee_id) = ?)
          AND (t.topic_id = ? OR t.topic_title IS NOT NULL)
        ORDER BY
          CASE WHEN t.topic_id = ? THEN 0 ELSE 1 END,
          t.created_at DESC
        """,
        (emp_id, emp_id, PRODUCT_TOPIC, PRODUCT_TOPIC),
    )
    if not tests:
        # Fallback: latest test for this employee code regardless of topic
        tests = store.query(
            """
            SELECT TOP 1 t.id
            FROM dbo.tests t
            WHERE t.employee_code = ?
            ORDER BY t.created_at DESC
            """,
            (emp_id,),
        )
    if not tests:
        return []
    rows = store.query(
        """
        SELECT question_text
        FROM dbo.test_questions
        WHERE test_id = ?
        ORDER BY question_index ASC
        """,
        (tests[0]["id"],),
    )
    return [str(r.get("question_text") or "").strip() for r in rows if str(r.get("question_text") or "").strip()]


def assigned_questions_for(employee_id: str) -> list[str]:
    live = _questions_from_active_test(employee_id)
    if live:
        return live
    norm = normalize_employee_id(employee_id)
    row = next((r for r in load_mapping_rows() if normalize_employee_id(r["employee_id"]) == norm), None)
    return list(row["assigned_questions"]) if row and row.get("assigned_questions") else []


def _domain_value(value) -> str:
    """Treat placeholder department 'general' as empty so mapping Domain can win."""
    text = str(value or "").strip()
    if not text or text.lower() == "general":
        return ""
    return text


def _eligible_accounts() -> list[dict]:
    rows = get_store().select("employees", {"product_qb_eligible": True},
                              ["employee_id", "email", "full_name", "department", "role", "product"], "employee_id asc")
    return [{
        "employee_id": str(r["employee_id"]), "full_name": r.get("full_name") or str(r["employee_id"]), "role": r.get("role") or "employee",
        "domain": _domain_value(r.get("department")), "product": r.get("product") or "", "email": r.get("email") or "", "ddh": "",
        "emp_status": "Confirmed", "remarks": "", "assigned_questions": [], "assigned_question_count": 0,
    } for r in rows if r.get("employee_id")]


def _merge(roster: dict, mapped: Optional[dict]) -> dict:
    if not mapped:
        return roster
    mapped_qs = list(mapped.get("assigned_questions") or [])
    roster_qs = list(roster.get("assigned_questions") or [])
    questions = mapped_qs or roster_qs
    # Prefer real mapping Domain/Product over placeholder roster values like "general".
    domain = _domain_value(mapped.get("domain")) or _domain_value(roster.get("domain"))
    product = (mapped.get("product") or "").strip() or (roster.get("product") or "").strip()
    return {
        "employee_id": roster["employee_id"] or mapped["employee_id"], "full_name": roster["full_name"] or mapped["full_name"],
        "role": roster["role"] or mapped["role"], "domain": domain,
        "product": product, "email": roster["email"] or mapped["email"], "ddh": roster["ddh"] or mapped["ddh"],
        "emp_status": mapped["emp_status"] or roster["emp_status"], "remarks": mapped["remarks"] or roster["remarks"],
        "assigned_questions": questions,
        "assigned_question_count": mapped.get("assigned_question_count")
        or roster.get("assigned_question_count")
        or len(questions),
    }


def manifest_from_results(results: list[dict]) -> dict[str, str]:
    out: dict[str, str] = {}
    for test in results:
        if test.get("topicId") == PRODUCT_TOPIC and test.get("employeeId"):
            out.setdefault(str(test["employeeId"]).strip(), test["id"])
    return out


def merge_portal_data(profile_rows: list[dict], results: list[dict], manifest: dict[str, str]) -> list[dict]:
    by_emp: dict[str, list[dict]] = {}
    for test in results:
        for raw in (test.get("employeeId"), test.get("employeeUuid"), test.get("employeeCode")):
            key = normalize_employee_id(raw)
            if not key:
                continue
            bucket = by_emp.setdefault(key, [])
            if not any(t["id"] == test["id"] for t in bucket):
                bucket.append(test)
    out = []
    for row in profile_rows:
        key = normalize_employee_id(row["employee_id"])
        tests = [t for t in by_emp.get(key, []) if not str(t.get("topicId") or "").startswith(HISTORY_PREFIX)]
        manifest_id = manifest.get(row["employee_id"]) or manifest.get(key)
        status_of = lambda t: str(t.get("status") or "").lower()
        completed = sorted((t for t in tests if status_of(t) == "completed"), key=lambda t: str(t.get("completedAt") or ""), reverse=True)
        assigned = (next((t for t in tests if t["id"] == manifest_id and status_of(t) == "completed"), None)
                    or next((t for t in tests if t.get("topicId") == PRODUCT_TOPIC and status_of(t) == "completed"), None)
                    or next((t for t in tests if t["id"] == manifest_id), None)
                    or next((t for t in tests if t.get("topicId") == PRODUCT_TOPIC), None))
        primary_completed = (next((t for t in completed if assigned and t["id"] == assigned["id"]), None)
                             or next((t for t in completed if t["id"] == manifest_id), None)
                             or next((t for t in completed if t.get("topicId") == PRODUCT_TOPIC), None)
                             or (completed[0] if completed else None))
        primary = primary_completed or assigned or (tests[0] if tests else None)
        score = (primary_completed.get("correctCount") if primary_completed.get("correctCount") is not None else primary_completed.get("score")) if primary_completed else None
        score_max = (primary or {}).get("totalQuestions") or row.get("assigned_question_count") or 25
        answered = (primary or {}).get("answeredCount") or 0
        total = (primary or {}).get("totalQuestions") or row.get("assigned_question_count") or 0
        completed_at = (primary_completed or {}).get("completedAt") or ((primary or {}).get("completedAt") if (primary or {}).get("status") == "completed" else None)
        raw_status = "completed" if primary_completed and primary_completed.get("status") == "completed" else ((assigned or {}).get("status") or (primary or {}).get("status"))
        test_id = (primary_completed or assigned or primary or {}).get("id") or manifest_id
        status_test_id = (assigned or primary or {}).get("id") or manifest_id
        if score is None and raw_status == "completed":
            score = (primary or {}).get("correctCount") if (primary or {}).get("correctCount") is not None else (primary or {}).get("score")
        out.append({
            **row, "product": format_product_display_name(row.get("product")), "test_id": test_id,
            "test_status": derive_status(row.get("assigned_question_count") or (25 if manifest_id else 0), status_test_id, raw_status, answered, total,
                                         (assigned or primary or {}).get("startedAt")),
            "score": score, "score_max": score_max, "completed_at": completed_at,
            "tests": [{"id": t["id"], "topicTitle": format_topic_title_for_display(t.get("topicTitle")), "subjectTitle": t.get("subjectTitle"),
                       "difficulty": t.get("difficulty"), "totalQuestions": t.get("totalQuestions"), "status": t.get("status"),
                       "score": t.get("correctCount") if t.get("correctCount") is not None else t.get("score"), "scoreMax": t.get("totalQuestions"),
                       "videoUrl": t.get("videoUrl"), "hasRecording": bool(t.get("hasRecording")), "proctoring": t.get("proctoring"),
                       "startedAt": t.get("startedAt"), "completedAt": t.get("completedAt")} for t in tests],
        })
    return out


def build_portal_employees(results: list[dict]) -> list[dict]:
    manifest = manifest_from_results(results)
    accounts = _eligible_accounts()
    mapping = load_mapping_rows()
    by_id = {normalize_employee_id(r["employee_id"]): r for r in mapping}
    counts = {}
    for test in results:
        if test.get("topicId") == PRODUCT_TOPIC:
            counts[normalize_employee_id(test.get("employeeId"))] = test.get("totalQuestions") or 25
    profiles: list[dict] = []
    for row in accounts:
        key = normalize_employee_id(row["employee_id"])
        merged = _merge(row, by_id.get(key))
        if not merged["assigned_question_count"]:
            merged["assigned_question_count"] = counts.get(key) or len(merged["assigned_questions"])
        profiles.append(merged)
    existing = {normalize_employee_id(r["employee_id"]) for r in profiles}
    for row in mapping:
        key = normalize_employee_id(row["employee_id"])
        if key not in existing:
            profiles.append(row)
            existing.add(key)
    merged = merge_portal_data(profiles, results, manifest)
    merged.sort(key=lambda r: (str(r.get("full_name") or "").lower(), str(r.get("employee_id") or "")))
    return merged
