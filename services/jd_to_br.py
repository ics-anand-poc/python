import io
import json
import logging
import re
from datetime import datetime, timezone
from typing import Optional

from services import llm, resumes

log = logging.getLogger("talentscope")

PROMPT = """
You are an expert HR Recruiting and AI Parsing system. Analyze the following Job Description text and extract structured information.

Provide your response in EXACT JSON format. Map and normalize skills. Detect duplicates. Do not include any markdown backticks or "json" prefix. The response must be a single, clean JSON object matching the schema below:

{
  "job_title": "Clean, official job title (e.g. L2 Production Support Engineer)",
  "department": "Department type if mentioned, e.g. Technical, Delivery (default to 'Technical')",
  "experience": "Experience range, e.g., 2-5 years",
  "support_level": "Identify Support Level: L1, L2, L3, L1/L2, etc. (default if not found: null)",
  "employment_type": "e.g., Full Time, Contract, Onsite, Hybrid",
  "shift_timing": "e.g., 24x7, Rotational, Day Shift, Irving TX Onsite",
  "skills": ["Array of technical skills such as SQL, PostgreSQL, Oracle, MySQL, Linux, Windows, Python, Shell scripting, APIs, Microservices, Kafka, Spring, Java, DevOps, JavaScript, React JS, Spring Boot, Cassandra DB, Network Security, CCNA etc."],
  "monitoring_tools": ["Array of tools like: Splunk, Dynatrace, AppDynamics, Datadog, Prometheus, Grafana, ELK"],
  "cloud_platforms": ["Array of platforms like: AWS, Azure, GCP"],
  "soft_skills": ["Array of soft skills like: Communication, Problem solving, Stakeholder management, Leadership, Teamwork"],
  "responsibilities": ["Array of core responsibilities like: Incident management, RCA, Monitoring, Troubleshooting, Deployment support, Batch job support, Release activities"],
  "tools": ["Array of tools like: Jira, ServiceNow, Remedy, Control-M, Autosys, Docker, Kubernetes, Nginx, Tomcat, Git"],
  "metrics": ["Array of service metrics like: SLA adherence, MTTR, Incident resolution rate"]
}

Normalize extracted items to standard capitalizations.

Job Description Text:
-------------------------------------------
__TEXT__
-------------------------------------------
"""

KEYWORDS = {
    "skills": ["sql", "postgresql", "oracle", "mysql", "linux", "windows", "unix", "python", "shell scripting", "apis", "microservices",
               "kafka", "java", "springboot", "javascript", "html", "css", "react", "angular", "hibernate", "spring boot"],
    "monitoring_tools": ["splunk", "dynatrace", "appdynamics", "datadog", "prometheus", "grafana", "elk"],
    "cloud_platforms": ["aws", "azure", "gcp", "gds"],
    "soft_skills": ["communication", "problem solving", "stakeholder management", "teamwork", "leadership"],
    "responsibilities": ["incident management", "rca", "root cause analysis", "monitoring", "troubleshooting", "deployment support",
                         "batch job support", "on-call rotations", "release activities"],
    "tools": ["jira", "servicenow", "remedy", "control-m", "autosys", "docker", "kubernetes", "nginx", "tomcat", "git"],
    "metrics": ["sla adherence", "mttr", "incident resolution rate", "sla"],
}
RENAMES = [("Sql", "SQL"), ("Apis", "APIs"), ("Rca", "RCA"), ("Mttr", "MTTR"), ("Sla", "SLA"), ("Gcp", "GCP"), ("Aws", "AWS"),
           ("Postgresql", "PostgreSQL"), ("Servicenow", "ServiceNow")]
TEXT_EXTS = ("docx", "pdf", "html", "htm", "txt")


def parse_document(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in TEXT_EXTS:
        raise ValueError(f"Unsupported file type: .{ext}. Only .docx, .pdf, .txt, and .html files are supported.")
    return resumes.extract_text(data, filename)


def keyword_fallback(text: str) -> dict:
    lower = text.lower()
    title = "Software Engineer"
    match = re.search(r"(Job Title|Title|Designation)\s*:\s*([^\r\n]+)", text, re.I)
    if match:
        title = match.group(2).strip()
    else:
        first = text.split("\n")[0].strip()
        if first and len(first) < 60:
            title = re.sub(r"^(Job Title|Title|Position|Role)\s*:\s*", "", first, flags=re.I)
    exp = re.search(r"(\d+\s*(?:-|to|\+|–|—)?\s*\d*)\s*(years|yrs|year|yr)", text, re.I)
    level = None
    if "l1" in lower and "l2" in lower:
        level = "L1/L2"
    elif "l3" in lower:
        level = "L3"
    elif "l2" in lower:
        level = "L2"
    elif "l1" in lower:
        level = "L1"
    out = {
        "job_title": title, "department": "Technical", "experience": exp.group(0).strip() if exp else "3+ years", "support_level": level,
        "employment_type": "Contract" if "contract" in lower else "Full Time",
        "shift_timing": "Rotational" if "rotational" in lower else "24x7" if "24x7" in lower else "Day Shift",
        **{k: [] for k in KEYWORDS},
    }
    for category, words in KEYWORDS.items():
        for word in words:
            if re.search(rf"\b{re.escape(word)}\b", lower, re.I):
                name = " ".join(w[:1].upper() + w[1:] for w in word.split(" "))
                for src, dst in RENAMES:
                    name = name.replace(src, dst, 1)
                out[category].append(name)
    out["skills"] = out["skills"] or ["SQL", "Unix"]
    out["soft_skills"] = out["soft_skills"] or ["Communication", "Problem solving"]
    out["responsibilities"] = out["responsibilities"] or ["Troubleshooting", "Monitoring"]
    out["tools"] = out["tools"] or ["Jira", "ServiceNow"]
    out["metrics"] = out["metrics"] or ["SLA adherence"]
    return out


def extract_details(raw_text: str) -> dict:
    try:
        result = llm.complete_json(PROMPT.replace("__TEXT__", raw_text[:12000]), temperature=0.1, max_tokens=2048, timeout_s=300)
        if isinstance(result, dict) and result.get("job_title"):
            return result
    except llm.LlmError as err:
        log.warning("Qwen JD extraction failed, using keyword fallback: %s", err)
    return keyword_fallback(raw_text)


def grade_from_experience(exp: Optional[str]) -> str:
    numbers = re.findall(r"\d+(?:\.\d+)?", exp or "")
    if not numbers:
        return "E2"
    years = float(numbers[0])
    for upper, grade in ((1, "E0"), (3, "E1"), (6, "E2"), (9, "E3"), (12, "E4")):
        if years < upper:
            return grade
    return "E5/E6"


def build_workbook(template: bytes, jd_files: list[tuple[str, bytes]], id_mapping: dict) -> bytes:
    from copy import copy

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(template))
    sheet = next((workbook[n] for n in ("BR _Raw Data", "Global TMH Demand an_21Oct_1715", "BR_Raw Data") if n in workbook.sheetnames),
                 workbook.worksheets[0])
    last = 1
    for r in range(sheet.max_row, 0, -1):
        if any(c.value not in (None, "") for c in sheet[r]):
            last = r
            break
    headers = {c.column: str(c.value).strip() for c in sheet[1] if c.value}
    current_num: Optional[int] = None
    suffix = ""
    for filename, data in jd_files:
        raw = parse_document(filename, data)
        if not raw.strip():
            continue
        details = extract_details(raw)
        template_row = last if last >= 2 else 2
        new_row = last + 1
        custom = id_mapping.get(filename) or ""
        auto_id = ""
        if custom:
            auto_id = custom
            if custom.isdigit():
                current_num, suffix = int(custom) + 1, ""
            else:
                current_num = None
        elif current_num is not None:
            auto_id = f"{current_num}{suffix}"
            current_num += 1
        else:
            col = next((c for c, h in headers.items() if h == "Auto req ID"), None)
            prev = sheet.cell(row=template_row, column=col).value if col else None
            if isinstance(prev, str):
                if prev.endswith("BR") and prev[:-2].isdigit():
                    auto_id, current_num, suffix = f"{int(prev[:-2]) + 1}BR", int(prev[:-2]) + 2, "BR"
                elif prev.isdigit():
                    auto_id, current_num, suffix = str(int(prev) + 1), int(prev) + 2, ""
            if not auto_id:
                import random

                auto_id = str(random.randint(40000, 49999))
                current_num, suffix = int(auto_id) + 1, ""
        skills = ", ".join(dict.fromkeys([*(details.get("skills") or []), *(details.get("monitoring_tools") or []),
                                          *(details.get("cloud_platforms") or [])]))
        mapping = {
            "Auto req ID": auto_id, "Current Req Status": "Open", "Grade": grade_from_experience(details.get("experience")),
            "Designation": details.get("job_title"), "Recruiter": "", "Department Type": "Technical", "BU": "ITS - TMH - Delivery",
            "Client Interview?": "Yes", "Mandatory Skills": skills, "Entity": "OFFSHORE", "Client Name": "IRON MOUNTAIN",
            "Billing Type": "Billable", "Project": "IM DXP-IDP 2025", "Requester ID": "1026374",
            "TAG Manager": "Antony, Nithin (1027544)", "RM Name": "Hippargi, Anil (1017237)", "Job description": raw[:5000],
            "Joining Location": "Bangalore - Global Axis", "Backfill for Employee Name": "",
            "Date Approved": datetime.now(timezone.utc).date().isoformat(), "No. of Positions": 1, "Positions Remaining": 1,
            "Sourcing Type": "External - India", "Requirement Type": "New",
            "ST (Bill Rate) Enter only numeric value and 0 for Non-Billable": 5.5,
        }
        for col, header in headers.items():
            cell = sheet.cell(row=new_row, column=col)
            cell.value = mapping.get(header, "")
            source = sheet.cell(row=template_row, column=col)
            if source.has_style:
                cell.font, cell.fill, cell.border = copy(source.font), copy(source.fill), copy(source.border)
                cell.alignment, cell.number_format = copy(source.alignment), source.number_format
        last = new_row
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
