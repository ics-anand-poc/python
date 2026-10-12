import csv
import hashlib
import io
import logging
import re
import zipfile
from datetime import datetime, timezone
from typing import Optional

from core.db import get_store
from core.security import safe_storage_filename
from services import corp_pool, dashboard_sync, deleted_requirements as deleted_req, docs_storage, jd_to_br, resumes, skill_match
from services.structured_log import write_log

log = logging.getLogger("talentscope")
JD_DOC_EXT = re.compile(r"\.(docx|doc|pdf|txt|html|htm)$", re.I)
CV_EXT = re.compile(r"\.(pdf|docx|doc|txt)$", re.I)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def cell_text(value) -> str:
    return "" if value is None else str(value).replace("\u00a0", " ").strip()


def normalize_br_id(raw: str) -> str:
    text = str(raw or "").strip()
    match = re.search(r"(\d+)\s*BR", text, re.I)
    if match:
        return f"{match.group(1)}BR"
    return f"{text}BR" if re.fullmatch(r"\d{4,}", text) else ""


def sanitize_corp_pool_file_name(name: str) -> str:
    base = re.split(r"[/\\]", str(name or ""))[-1] or "resume"
    cleaned = re.sub(r"_+", "_", re.sub(r"\s+", " ", re.sub(r"[^\w.\- ()\[\]]+", "_", base.replace("\u00a0", " ")))).strip() or "resume"
    if len(cleaned) <= 180:
        return cleaned
    ext = cleaned[cleaned.rfind("."):] if "." in cleaned else ""
    return cleaned[: max(1, 180 - len(ext))] + ext


def is_corp_pool_roster_file_name(name: str) -> bool:
    n = re.sub(r"[_'`’]", " ", str(name or "").lower())
    return bool(re.search(r"\.(xlsx|xls|csv)$", n)) and any(k in n for k in ("corp pool", "active list", "employee list"))


EMP_HEADER_HINTS = ("emp no", "emp_no", "employee id", "emp id", "emp name", "employee name")


def _is_emp_header_cell(value) -> bool:
    text = cell_text(value).lower()
    return any(h in text for h in EMP_HEADER_HINTS)


def excel_looks_like_corp_pool_roster(data: bytes) -> bool:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(min_row=1, max_row=10, values_only=True):
                if any(_is_emp_header_cell(v) for v in row):
                    return True
    except Exception:
        return False
    return False


def _norm_header(value) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\.+$", "", re.sub(r"[_-]", " ", cell_text(value).lower()))).strip()


def _find_col(headers: list[str], names: list[str], loose: bool) -> int:
    for name in names:
        norm = re.sub(r"[_-]", " ", name.strip().lower())
        for i, h in enumerate(headers):
            if not h:
                continue
            if len(norm) <= 3:
                if h == norm:
                    return i
            elif h == norm or norm in h or (loose and h in norm):
                return i
    return -1


def _columns(headers: list[str], loose: bool, with_id_alias: bool, with_skill_alias: bool, with_mail_alias: bool) -> dict:
    return {
        "id": _find_col(headers, ["emp no", "employee id", "emp id", "employee code", *(["id"] if with_id_alias else [])], loose),
        "name": _find_col(headers, ["emp name", "employee name", "full name", "name"], loose),
        "dept": _find_col(headers, ["business unit", "sbu", "bu", "department", "dept"], loose),
        "skills": _find_col(headers, ["detailed skills", "skills bucket", "top 3 skills", *(["primary skill"] if with_skill_alias else []), "skills"], loose),
        "status": _find_col(headers, ["status", "availability"], loose),
        "grade": _find_col(headers, ["grade", "level"], loose),
        "mail": _find_col(headers, ["official mail id", "official email", "email", "mail id", *(["mail"] if with_mail_alias else [])], loose),
        "role": _find_col(headers, ["designation", "role", "position"], loose),
    }


def _fallback_id(file: str, name: str, row: int) -> str:
    return "EMP" + hashlib.md5(f"{file}:{name}:{row}".encode("utf-8"), usedforsecurity=False).hexdigest()[:10]


def _person(row, cols: dict, file: str, row_index: int, defaults_empty_ok: bool) -> Optional[dict]:
    def at(key: str) -> str:
        idx = cols[key]
        return cell_text(row[idx]) if idx != -1 and idx < len(row) else ""

    emp_no, emp_name = at("id"), at("name")
    if not emp_no and not emp_name:
        return None
    employee_id = emp_no or _fallback_id(file, emp_name, row_index)
    department = at("dept") or "Engineering"
    status = at("status") or "Active"
    grade = at("grade") or "E1"
    designation = at("role") or "Support Engineer"
    return {"employee_id": employee_id, "full_name": emp_name or "Unknown Employee", "email": at("mail") or f"{employee_id}@example.com",
            "department": department, "skills": at("skills"), "grade": grade, "designation": designation, "status": status,
            "shortlisted": False, "score": 0, "matchingSkills": [], "source_file": file}


def parse_roster_excel(data: bytes, file: str) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), data_only=True)
    rows: list[list] = []
    header_idx = -1
    for sheet in workbook.worksheets:
        table = [list(r) for r in sheet.iter_rows(values_only=True)]
        found = next((i for i, r in enumerate(table[:10]) if any(_is_emp_header_cell(v) for v in r)), -1)
        if found != -1:
            rows, header_idx = table, found
            break
    if header_idx == -1 and workbook.worksheets:
        rows = [list(r) for r in workbook.worksheets[0].iter_rows(values_only=True)]
        header_idx = 0
    if len(rows) <= header_idx + 1:
        return []
    cols = _columns([_norm_header(v) for v in rows[header_idx]], False, False, False, False)
    out = []
    for index in range(header_idx + 1, len(rows)):
        person = _person(rows[index], cols, file, index, False)
        if person:
            out.append(person)
    return out


def _decode_text(data: bytes) -> str:
    if data[:2] == b"\xff\xfe":
        return data[2:].decode("utf-16-le", errors="replace")
    if data[:2] == b"\xfe\xff":
        return data[2:].decode("utf-16-be", errors="replace")
    return data.decode("utf-8-sig", errors="replace").replace("\r\n", "\n").replace("\r", "\n")


def parse_roster_csv(data: bytes, file: str) -> list[dict]:
    lines = [ln.rstrip() for ln in _decode_text(data).split("\n") if ln.strip()]
    if len(lines) < 2:
        return []
    counts = {d: lines[0].count(d) for d in (",", ";", "\t")}
    delimiter = max(counts, key=lambda d: counts[d]) if max(counts.values()) > 0 else ","
    rows = list(csv.reader(lines, delimiter=delimiter))
    cols = _columns([re.sub(r"[_-]", " ", str(h or "").strip().lower()) for h in rows[0]], True, True, True, True)
    out = []
    for index, row in enumerate(rows[1:], 1):
        if not row or all(not str(c or "").strip() for c in row):
            continue
        person = _person(row, cols, file, index, True)
        if person:
            out.append(person)
    return out


PHONE_DIGITS = (10, 11, 12)


def _plausible_emp_id(value: str) -> bool:
    ident = str(value or "").strip()
    if not re.fullmatch(r"[A-Za-z]?\d{4,12}", ident):
        return False
    digits = re.sub(r"\D", "", ident)
    if len(digits) in PHONE_DIGITS:
        return False
    return not (len(digits) == 4 and 1970 <= int(digits) <= 2035)


def extract_employee_id_from_cv(text: str, file: str) -> str:
    normalized = str(text or "").replace("\u00a0", " ")
    for pattern in (r"employee\s*(?:id|code|number|no)\s*[:#.\-|]*\s*([A-Za-z]?\d{4,12})\b",
                    r"emp(?:loyee)?\s*(?:id|no|code|number)\s*[:#.\-|]*\s*([A-Za-z]?\d{4,12})\b",
                    r"staff\s*(?:id|code|no)\s*[:#.\-|]*\s*([A-Za-z]?\d{4,12})\b"):
        m = re.search(pattern, normalized, re.I)
        if m and _plausible_emp_id(m.group(1)):
            return m.group(1).strip()
    lines = [ln.strip() for ln in normalized.splitlines() if ln.strip()]
    for i in range(min(len(lines) - 1, 20)):
        if re.fullmatch(r"(employee\s*(?:id|code|number|no)|emp(?:loyee)?\s*(?:id|no)|staff\s*id)\s*[:#.\-]*", lines[i], re.I):
            m = re.match(r"([A-Za-z]?\d{4,12})\b", lines[i + 1])
            if m and _plausible_emp_id(m.group(1)):
                return m.group(1)
    m = re.search(r"employee\s*id[\s\S]{0,120}?(\d{5,10})", normalized, re.I)
    if m and _plausible_emp_id(m.group(1)):
        return m.group(1)
    m = re.search(r"\b([A-Za-z]?\d{5,10})\b", str(file or ""))
    if m and _plausible_emp_id(m.group(1)):
        return m.group(1)
    for line in lines[:12]:
        m = re.fullmatch(r"([A-Za-z]?\d{5,8})(?:\s+[A-Za-z].*)?", line)
        if m and _plausible_emp_id(m.group(1)):
            return m.group(1)
    return ""


CV_NAME_BLOCK = re.compile(r"^(career objectives?|objective|summary|professional summary|profile|experience|work experience|education|skills|technical skills|"
                           r"contact|contacts|declaration|projects|certifications?|about me|resume|curriculum vitae|personal details|employment history|"
                           r"key skills|highlights|achievements?)$", re.I)
CV_SKILLS = ["typescript", "javascript", "python", "java", "sql", "ms sql", "mysql", "postgresql", "windows", "linux", "unix", "protractor", "selenium",
             "cypress", "playwright", "jira", "git", "jenkins", "agile", "scrum", "waterfall", "stlc", "sdlc", "html", "css", "react", "angular", "node.js",
             "aws", "azure", "docker", "rest", "api", "postman", "manual testing", "automation testing"]


def _strip_leading_id(value: str) -> tuple[str, str]:
    m = re.match(r"^([A-Za-z]?\d{5,12})\s+(.+)$", str(value or "").strip())
    if m and _plausible_emp_id(m.group(1)):
        return m.group(1), m.group(2).strip()
    return "", str(value or "").strip()


def looks_like_person_name(line: str) -> bool:
    raw = re.sub(r"\s+", " ", str(line or "")).strip()
    text = re.sub(r"\s+", " ", re.sub(r"[!|]+", " ", _strip_leading_id(raw)[1])).strip()
    if not text or CV_NAME_BLOCK.match(text) or len(text) < 3 or len(text) > 50:
        return False
    if re.search(r"@|https?:|www\.", text, re.I) or "|" in raw:
        return False
    if re.match(r"^(successfully|experienced|worked|developed|led|responsible|managed|supporting)\b", text, re.I):
        return False
    words = [w for w in text.split(" ") if w]
    if not words or len(words) > 5 or (len(words) == 1 and len(words[0]) < 3):
        return False
    if any(re.fullmatch(r"20\d{2}", w) or re.fullmatch(r"infinite", w, re.I) for w in words):
        return False
    if not re.fullmatch(r"[A-Za-z][A-Za-z .'-]*", text):
        return False
    title = [w for w in words if re.fullmatch(r"[A-Z][a-zA-Z'.-]*", w) or re.fullmatch(r"[A-Z]\.?", w)]
    return len(title) >= -(-len(words) // 2)


def skills_from_cv_text(text: str) -> list[str]:
    lower = f" {str(text or '').lower()} "
    found = []
    for skill in CV_SKILLS:
        if re.search(rf"[^a-z0-9]{re.escape(skill)}[^a-z0-9]", lower, re.I) and skill not in found:
            found.append(skill)
    return found


def _labeled(text: str, labels: list[str]) -> str:
    joined = "|".join(re.sub(r"\s+", r"\\s*", re.escape(label).replace("\\ ", " ")) for label in labels)
    m = re.search(rf"(?:^|[\n\r])\s*(?:{joined})\s*[:|#]\s*([^\n\r]{{2,80}})", text, re.I)
    return m.group(1).strip() if m else ""


def corp_pool_profile_from_cv(file: str, text: str) -> dict:
    base = re.sub(r"\.[^/.]+$", "", file)
    m = re.search(r"\b(SDET|QA|Quality\s*Analyst|Developer|Engineer|Lead|Manager|Architect|Analyst|Consultant|Tester)\b", base, re.I)
    designation_file = m.group(0) if m else ""
    from_file_raw = re.sub(r"\s+", " ", re.sub(r"\b20\d{2}\b", "", re.sub(r"\b(SDET|QA|resume|cv|curriculum vitae|infinite)\b", "",
                           re.sub(r"\b\d+\s*(yoe|yrs?|years?)\b", "", re.sub(r"[_-]+", " ", base), flags=re.I), flags=re.I))).strip()
    file_id, file_rest = _strip_leading_id(from_file_raw)
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    labeled_id, labeled_rest = _strip_leading_id(re.sub(r"employee\s*id.*", "", _labeled(text, ["name", "candidate name", "employee name"]), flags=re.I))
    name_line = next((ln for ln in lines[:15] if looks_like_person_name(ln)), None)
    line_id, line_rest = _strip_leading_id(name_line) if name_line else ("", "")
    labeled_title = _labeled(text, ["title", "designation", "role", "position"])
    title_line = next((ln for ln in lines if len(ln) < 60 and "|" not in ln
                       and re.search(r"\b(SDET|engineer|developer|lead|manager|analyst|architect|tester)\b", ln, re.I)
                       and not re.search(r"employee\s*id|years? of|successfully|interoperability", ln, re.I) and not looks_like_person_name(ln)), "")
    mail_raw = _labeled(text, ["email id", "email", "mail id", "e-mail"]) or (re.search(r"[\w.-]+@[\w.-]+\.\w+", text) or [""])[0]
    mail = (re.search(r"[\w.-]+@[\w.-]+\.\w+", mail_raw) or [""])[0]
    employee_id = extract_employee_id_from_cv(text, file) or labeled_id or line_id or file_id
    name = (labeled_rest if looks_like_person_name(labeled_rest) else "") or line_rest or (file_rest if looks_like_person_name(file_rest) else "") or "Unknown"
    return {"name": name, "designation": labeled_title or designation_file or title_line or "Engineer", "email": mail, "employeeId": employee_id}


def _jd_for_scoring(active_jd_id: Optional[str]) -> tuple[Optional[str], str]:
    store = get_store()
    if not active_jd_id or active_jd_id == "all":
        row = store.first("job_descriptions", columns=["id", "jd_text"], order="created_at desc")
    else:
        row = store.first("job_descriptions", {"id": active_jd_id}, ["id", "jd_text"]) if deleted_req.UUID_OK(active_jd_id) else None
    return (row["id"], row.get("jd_text") or "") if row else (None, "")


def _score(jd_text: str, people: list[dict]) -> list[dict]:
    if not jd_text.strip() or not people:
        return [{"score": 0, "matchingSkills": []} for _ in people]
    try:
        return skill_match.score_people(jd_text, people)
    except skill_match.SkillMatchError as err:
        write_log("employee", "SKILL_MATCH_UNAVAILABLE", "failed", str(err))
        return [{"score": 0, "matchingSkills": []} for _ in people]


def _expand_zip(file: str, data: bytes, buffers: dict) -> list[str]:
    names, used, extracted, skipped = [], set(), 0, 0
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for info in archive.infolist():
                name = info.filename.replace("\\", "/")
                base = name.split("/")[-1]
                if info.is_dir() or name.startswith("__MACOSX") or any(p.startswith(".") for p in name.split("/")) \
                        or not re.search(r"\.(pdf|docx|doc|txt|csv|xlsx|xls)$", base, re.I) or info.file_size == 0 or info.file_size > 25 * 1024 * 1024:
                    skipped += 1
                    continue
                payload = archive.read(info)
                stored = _unique_name(base, used)
                try:
                    docs_storage.write_file("Corp Pool", stored, payload)
                except Exception as err:
                    write_log("employee", "UNZIP_STORE_FAILED", "failed", f"Parsed {stored} in memory after storage failed: {err}")
                buffers[stored] = payload
                names.append(stored)
                extracted += 1
        write_log("employee", "UNZIPPED_CORP_POOL", "success" if extracted else "failed",
                  f"Extracted {extracted} file(s) from {file}{f' (skipped {skipped})' if skipped else ''}" if extracted
                  else f"ZIP {file} had no resume/Excel/CSV files inside")
    except zipfile.BadZipFile as err:
        write_log("employee", "UNZIP_CORP_POOL_FAILED", "failed", f"Failed reading ZIP {file}: {err}")
    return names


def _unique_name(name: str, used: set) -> str:
    safe = sanitize_corp_pool_file_name(name)
    if safe.lower() not in used:
        used.add(safe.lower())
        return safe
    stem, ext = (safe[: safe.rfind(".")], safe[safe.rfind("."):]) if "." in safe else (safe, "")
    i = 2
    while f"{stem}_{i}{ext}".lower() in used:
        i += 1
    used.add(f"{stem}_{i}{ext}".lower())
    return f"{stem}_{i}{ext}"


def reclaim_misfiled_rosters() -> list[str]:
    moved = []
    for file in docs_storage.list_files("Resumes"):
        if not is_corp_pool_roster_file_name(file):
            continue
        data = docs_storage.read_file("Resumes", file)
        if data is None:
            continue
        stored = sanitize_corp_pool_file_name(file)
        docs_storage.write_file("Corp Pool", stored, data)
        docs_storage.delete_file("Resumes", file)
        moved.append(stored)
        write_log("employee", "RECLAIMED_CORP_POOL_ROSTER", "success", f"Moved {file} from Resumes to Corp Pool as {stored}")
    return moved


def refresh_employees(active_jd_id: Optional[str] = None, incoming: Optional[list[tuple[str, bytes]]] = None) -> dict:
    reclaimed = reclaim_misfiled_rosters()
    buffers: dict[str, bytes] = {}
    keys: set[str] = set()
    for name, data in incoming or []:
        stored = sanitize_corp_pool_file_name(name)
        buffers[stored] = data
        keys.add(stored.lower())
    if keys:
        keys.update(sanitize_corp_pool_file_name(f).lower() for f in reclaimed)
    files = docs_storage.list_files("Corp Pool")
    if keys:
        files = [f for f in files if sanitize_corp_pool_file_name(f).lower() in keys]
    for stored in buffers:
        if not any(f.lower() == stored.lower() for f in files):
            files.append(stored)
    expanded: list[str] = []
    for file in files:
        if file.lower().endswith(".zip"):
            expanded.extend(_expand_zip(file, buffers.get(file) or docs_storage.read_file("Corp Pool", file) or b"", buffers))
        else:
            expanded.append(file)
    files = list(dict.fromkeys(expanded))
    if keys and not files:
        raise ValueError("The ZIP had no resume PDF/DOCX or employee Excel/CSV files inside. Put those files in the ZIP and upload again."
                         if any(k.endswith(".zip") for k in keys) else
                         "The uploaded Corp Pool file could not be read after storage. Try renaming it (avoid apostrophes) and upload again.")

    def data_for(file: str) -> bytes:
        return buffers.get(file) or docs_storage.read_file("Corp Pool", file) or b""

    parsed: list[dict] = []
    for file in files:
        low = file.lower()
        data = data_for(file)
        try:
            if low.endswith(".csv"):
                parsed.extend(parse_roster_csv(data, file))
            elif low.endswith((".xlsx", ".xls")) or (data[:2] == b"PK" and not CV_EXT.search(low)):
                people = parse_roster_excel(data, file)
                parsed.extend(people)
                write_log("employee", "PARSED_CORP_POOL_EXCEL", "success", f"Parsed {len(people)} people from {file}")
            elif CV_EXT.search(low):
                if not data:
                    write_log("employee", "PARSE_CV_EMPTY", "failed", f"Corp Pool CV {file} is empty")
                    continue
                text = resumes.extract_text(data, file).strip()
                if not text:
                    write_log("employee", "PARSE_CV_EMPTY", "failed", f"No text extracted from Corp Pool CV {file}")
                    continue
                profile = corp_pool_profile_from_cv(file, text)
                cv_skills = skills_from_cv_text(text)
                employee_id = profile["employeeId"] or "CV" + hashlib.md5(file.lower().encode("utf-8"), usedforsecurity=False).hexdigest()[:10]
                parsed.append({"employee_id": employee_id, "full_name": profile["name"], "email": profile["email"] or f"{employee_id}@corp-pool.local",
                               "department": "Engineering", "skills": ", ".join(cv_skills), "grade": "", "designation": profile["designation"],
                               "status": "Active", "shortlisted": False, "score": 0, "matchingSkills": cv_skills, "source_file": file,
                               "_cv_text": text})
                write_log("employee", "PARSED_CORP_POOL_CV", "success", f"Added {profile['name']} to Corp Pool from {file}")
        except Exception as err:
            write_log("employee", "PARSE_CORP_POOL_FAILED", "failed", f"Error parsing Corp Pool file {file}: {err}")

    seen, unique = set(), []
    for emp in parsed:
        if emp["employee_id"] not in seen:
            seen.add(emp["employee_id"])
            unique.append(emp)
    _, jd_text = _jd_for_scoring(active_jd_id)
    scoring = [{**e, "skills": e.pop("_cv_text")} if "_cv_text" in e else e for e in unique]
    for emp, res in zip(unique, _score(jd_text, scoring)):
        emp["score"] = res["score"]
        if res.get("matchingSkills"):
            emp["matchingSkills"] = res["matchingSkills"]
    incoming_upload = bool(keys)
    if incoming_upload:
        corp_pool.unmark_deleted([e["employee_id"] for e in unique], [*keys, *files])
    deleted = corp_pool.load_deleted()
    if not incoming_upload:
        unique = [e for e in unique if not corp_pool.is_deleted(deleted, e["employee_id"])]
    existing = [e for e in corp_pool.load_roster() if not corp_pool.is_deleted(deleted, e["employee_id"])]
    if not unique:
        if incoming_upload:
            write_log("employee", "INCOMING_CORP_POOL_EMPTY", "failed", "Uploaded Corp Pool file produced 0 people; left the existing list unchanged")
            raise ValueError("The file was stored, but nobody was added to Corp Pool. Use a readable resume PDF/DOCX, an employee Excel/CSV, or a ZIP of those files.")
        write_log("employee", "SKIP_EMPTY_CORP_POOL", "success", "Skipped empty Corp Pool refresh to preserve Employee Portal roster and tests")
        return {"success": True, "loaded": len(existing), "added": 0, "updated": 0, "skippedDeleted": 0, "employees": existing}
    by_id = {str(e["employee_id"]): e for e in existing}
    by_email = {str(e.get("email") or "").strip().lower(): str(e["employee_id"]) for e in existing if str(e.get("email") or "").strip()}
    added = updated = 0
    batch = _now()
    for item in unique:
        email = str(item.get("email") or "").strip().lower()
        existing_id = item["employee_id"] if item["employee_id"] in by_id else by_email.get(email) if email else None
        if existing_id and existing_id in by_id:
            prev = by_id[existing_id]
            generated = lambda v: bool(re.fullmatch(r"CV[a-f0-9]{8,}", str(v or "").strip(), re.I))
            keep = item["employee_id"] if generated(existing_id) and not generated(item["employee_id"]) else existing_id
            if keep != existing_id:
                by_id.pop(existing_id, None)
            merged = {**item, "employee_id": keep, "shortlisted": prev.get("shortlisted"), "score_override": prev.get("score_override"),
                      "score_override_jd_id": prev.get("score_override_jd_id"), "manually_edited": prev.get("manually_edited"),
                      "uploaded_at": prev.get("uploaded_at") or (batch if incoming_upload else None),
                      "upload_batch": batch if incoming_upload else prev.get("upload_batch")}
            if prev.get("manually_edited"):
                for field in ("full_name", "email", "department", "skills", "designation", "grade", "product", "status"):
                    merged[field] = prev.get(field)
            merged["score"] = prev["score_override"] if isinstance(prev.get("score_override"), (int, float)) else item["score"]
            by_id[keep] = {k: v for k, v in merged.items() if v is not None}
            if email:
                by_email[email] = keep
            updated += 1
        else:
            by_id[item["employee_id"]] = {**item, "uploaded_at": batch, "upload_batch": batch}
            if email:
                by_email[email] = item["employee_id"]
            added += 1
    final = [e for e in by_id.values() if not corp_pool.is_deleted(deleted, e["employee_id"])]
    corp_pool.save_roster(final)
    write_log("employee", "SYNC_EMPLOYEE_POOL", "success",
              f"Corp Pool now has {len(final)} people (added {added}, updated {updated})" if incoming_upload
              else f"Successfully loaded {len(final)} employees from Corp Pool uploads")
    return {"success": True, "loaded": len(final), "added": added, "updated": updated, "skippedDeleted": 0, "employees": final}


def compose_requirement_text(designation: str, skills: str, jd_body: str) -> str:
    parts = []
    if designation:
        parts.append(f"Job Title: {designation}")
    if skills:
        parts.append(f"Mandatory Skills: {skills}")
    if jd_body:
        parts.append(jd_body)
    return "\n\n".join(parts).strip()


def parse_br_workbook(data: bytes, source: str) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for sheet in workbook.worksheets:
        name = (sheet.title or "").lower()
        if not name.strip() or "pivot" in name or "summary" in name:
            continue
        table = [list(r) for r in sheet.iter_rows(values_only=True)]
        header_idx = next((i for i, r in enumerate(table[:10]) if any(v and ("auto req id" in str(v).strip().lower() or "br id" in str(v).strip().lower()) for v in r)), -1)
        if header_idx == -1:
            continue
        header = [re.sub(r"[_-]", " ", str(v).strip().lower()) if v else "" for v in table[header_idx]]

        def col(names: list[str]) -> int:
            for n in names:
                if n in header:
                    return header.index(n)
            return -1

        id_i, title_i = col(["auto req id", "br id", "id"]), col(["designation", "job title", "role", "position"])
        skills_i, jd_i = col(["mandatory skills", "skills", "detailed skills"]), col(["job description", "jd"])
        rm_i = col(["rm email", "recruiter", "requester id", "rm name", "reporting manager"])
        if id_i == -1:
            continue
        for row in table[header_idx + 1:]:
            at = lambda i: cell_text(row[i]) if i != -1 and i < len(row) else ""
            auto_id = normalize_br_id(at(id_i))
            if not auto_id:
                continue
            designation, skills, body, rm = at(title_i), at(skills_i), at(jd_i), at(rm_i)
            composed = compose_requirement_text(designation or "Technical Role", skills, body)
            if composed:
                out.append({"autoReqId": auto_id, "designation": designation or "Technical Role", "skills": skills, "jdBody": body,
                            "rmEmail": rm.lower() if "@" in rm else "admin@infinite.com", "sourceFile": source, "composedText": composed})
    return out


def _weight(row: dict) -> int:
    return len(row["composedText"]) + (500 if len(row["jdBody"]) > 80 else 0)


def merge_br_requirements(rows: list[dict]) -> list[dict]:
    by_id: dict[str, dict] = {}
    for row in rows:
        prev = by_id.get(row["autoReqId"])
        if not prev or _weight(row) > _weight(prev):
            by_id[row["autoReqId"]] = row
    groups: dict[str, list] = {}
    for row in by_id.values():
        text = row["composedText"].strip().lower()
        groups.setdefault(f"id:{row['autoReqId']}" if len(text) < 40 else text, []).append(row)
    merged = []
    for items in groups.values():
        merged.extend(sorted(items, key=_weight, reverse=True)[:2])
    return merged


def _real_rm(value) -> bool:
    email = str(value or "").strip().lower()
    return "@" in email and email != "admin@infinite.com"


def prefer_rm(existing, incoming, actor, prefer_actor=False) -> str:
    norm = lambda v: str(v or "").strip().lower()
    if not prefer_actor and _real_rm(existing):
        return norm(existing)
    if _real_rm(incoming):
        return norm(incoming)
    if _real_rm(actor) and (prefer_actor or not norm(existing)):
        return norm(actor)
    if _real_rm(existing):
        return norm(existing)
    return norm(existing) or norm(incoming) or "admin@infinite.com"


def refresh_requirements(incoming_br: Optional[list[tuple[str, bytes]]] = None, incoming_jd: Optional[list[tuple[str, bytes]]] = None,
                         actor_email: str = "") -> dict:
    store = get_store()
    actor = (actor_email or "").strip().lower()
    br_incoming = {n.lower(): d for n, d in incoming_br or []}
    jd_incoming = {n.lower(): d for n, d in incoming_jd or []}
    deleted = deleted_req.load()
    now = _now()
    existing = {str(r["id"]).lower(): r for r in store.select("job_descriptions")}
    br_files = {f.lower(): f for f in docs_storage.list_files("BR") if re.search(r"\.(xlsx|xls)$", f, re.I)}
    for name in br_incoming:
        br_files.setdefault(name, name)
    jd_files = {f.lower(): f for f in docs_storage.list_files("JD") if JD_DOC_EXT.search(f)}
    for name in jd_incoming:
        if JD_DOC_EXT.search(name):
            jd_files.setdefault(name, name)
    rows_to_save: dict[str, dict] = {}
    incoming_rows = converted = 0
    parsed_all: list[dict] = []
    for key, file in br_files.items():
        is_incoming = key in br_incoming
        if not is_incoming:
            docs_storage.delete_file("BR", file)
            write_log("requirements", "REMOVED_EXTRA_BR_FILE", "success", f"Removed leftover BR file {file}")
            continue
        try:
            data = br_incoming.get(key) or docs_storage.read_file("BR", file) or b""
            parsed = parse_br_workbook(data, file)
            incoming_rows += len(parsed)
            parsed_all.extend(parsed)
            docs_storage.delete_file("BR", file)
            write_log("requirements", "MERGED_BR_INTO_REQUIREMENTS", "success", f"Read {len(parsed)} BR row(s) from {file}")
        except Exception as err:
            write_log("requirements", "PARSE_BR_FILE_FAILED", "failed", f"Failed parsing BR file {file}: {err}")
    keep: set[str] = set()
    for row in merge_br_requirements(parsed_all):
        uid = deleted_req.br_id_to_uuid(row["autoReqId"])
        file_name = f"{row['autoReqId']} | {row['sourceFile']}"
        deleted_req.unmark_deleted([{"id": uid, "brId": row["autoReqId"], "fileName": file_name}])
        keep.add(uid)
        prior = existing.get(uid)
        keep_doc = bool(prior and JD_DOC_EXT.search(str(prior.get("file_name") or "")))
        rows_to_save[uid] = {
            "id": uid, "jd_text": prior["jd_text"] if keep_doc and prior.get("jd_text") else row["composedText"],
            "rm_email": prefer_rm(prior.get("rm_email") if prior else "", row["rmEmail"], actor, not prior),
            "file_name": prior["file_name"] if keep_doc else file_name, "created_at": (prior or {}).get("created_at") or now}
    labels = [str(r.get("file_name") or "") for r in existing.values()]
    for key, file in jd_files.items():
        try:
            data = jd_incoming.get(key) or docs_storage.read_file("JD", file) or b""
            text = resumes.extract_text(data, file)
            if not text.strip():
                write_log("requirements", "CONVERT_JD_EMPTY", "failed", f"No text extracted from JD {file}")
                continue
            linked = next((r for r in existing.values() if file.lower() in str(r.get("file_name") or "").lower()), None)
            br_id = normalize_br_id(file) or (normalize_br_id(str(linked.get("file_name") or "").split("|")[0]) if linked else "")
            if not br_id:
                if linked:
                    continue
                br_id = deleted_req.next_synthetic_br_id([*labels, *[f"{r['file_name']}" for r in rows_to_save.values()]])
            uid = deleted_req.br_id_to_uuid(br_id)
            is_incoming = key in jd_incoming
            parts = {"id": uid, "brId": br_id, "fileName": f"{br_id} | {file}"}
            if is_incoming:
                deleted_req.unmark_deleted([parts])
            elif deleted_req.is_deleted(deleted_req.load(), uid, parts["fileName"], br_id) or linked:
                continue
            details = jd_to_br.keyword_fallback(text)
            skills = ", ".join(dict.fromkeys([*(details.get("skills") or []), *(details.get("tools") or [])]))
            prior = existing.get(uid)
            rows_to_save[uid] = {
                "id": uid, "jd_text": compose_requirement_text(details.get("job_title") or "Technical Role", skills, text) or text,
                "rm_email": prefer_rm((prior or {}).get("rm_email"), "", actor, is_incoming), "file_name": parts["fileName"],
                "created_at": (prior or {}).get("created_at") or now}
            keep.add(uid)
            converted += 1
            labels.append(parts["fileName"])
            write_log("requirements", "CONVERTED_JD_TO_BR", "success", f"Converted JD {file} into requirement {br_id}")
        except Exception as err:
            write_log("requirements", "CONVERT_JD_FAILED", "failed", f"Error converting JD {file}: {err}")
    live_deleted = deleted_req.load()
    processed = 0
    for uid, row in rows_to_save.items():
        if uid not in keep and deleted_req.is_deleted(live_deleted, uid, row["file_name"]):
            continue
        store.upsert("job_descriptions", row, "id")
        processed += 1
    dashboard_sync.bump_sync("refresh_requirements")
    return {"success": True, "processedBRs": processed, "convertedJDs": converted, "incomingBrRows": incoming_rows}


def refresh_candidates(active_jd_id: Optional[str] = None) -> dict:
    jd_id, jd_text = _jd_for_scoring(active_jd_id)
    processed = duplicates = 0
    for file in (f for f in docs_storage.list_files("Resumes") if re.search(r"\.(pdf|docx|doc)$", f, re.I)):
        try:
            data = docs_storage.read_file("Resumes", file) or b""
            file_hash = hashlib.sha256(data).hexdigest()
            prior = get_store().first("resumes", {"file_hash": file_hash}, ["id", "report"])
            if prior:
                report = prior.get("report")
                report = __import__("json").loads(report) if isinstance(report, str) and report else (report or {})
                if report.get("jdId") == jd_id:
                    duplicates += 1
                    continue
            result = resumes.process_resume(file, data, jd_text, jd_id, "admin@infinite.com", False, wait=True)
            processed += 1
            score = (result.get("report") or {}).get("jdMatchScore") or (result.get("analysis") or {}).get("overallScore") or 0
            name = (result.get("parsed") or {}).get("personal", {}).get("fullName") or re.sub(r"\.[^/.]+$", "", file)
            write_log("candidate-processing", f"SCREENED_{'SUITABLE' if score >= 40 else 'UNSUITABLE'}_CANDIDATE", "success",
                      f"Candidate {name} matches {score}% for JD {jd_id or 'latest'}")
        except Exception as err:
            write_log("candidate-processing", "SCREENING_FAILED", "failed", f"Error screening CV {file}: {err}")
    return {"success": True, "processed": processed, "duplicates": duplicates}


def refresh_interviews() -> dict:
    try:
        attempts = get_store().select("interview_attempts", columns=["resume_id"])
        count = len({a["resume_id"] for a in attempts if a.get("resume_id")})
        write_log("interview", "REFRESH_INTERVIEWS", "success", f"Interview data is live in the database ({count} candidate(s) with answers)")
        return {"success": True, "count": count}
    except Exception as err:
        write_log("interview", "REFRESH_INTERVIEWS_FAILED", "failed", f"Failed to read interview data: {err}")
        return {"success": False, "count": 0}
