import io
import logging
import re
from typing import Optional

from services import llm

log = logging.getLogger("talentscope")

KIND_TO_CATEGORY = {"corp_pool": "employee", "jd": "jd", "br": "br", "portal_mapping": "portal-mapping", "unknown": "resume"}
KIND_TO_PLACE = {
    "corp_pool": "Corp Pool (people list)", "jd": "Requirements JD/BR", "br": "Requirements JD/BR (BR workbook)",
    "portal_mapping": "Employee Portal mapping", "unknown": "Do not auto-place — pick the tab yourself",
}


def _norm(name: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[_'`’]", " ", str(name or "").lower())).strip()


def _token(line: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[_-]+", " ", line.lower())).strip()


def _find_column(preview: str, names: list[str]) -> str:
    for line in str(preview or "").split("\n")[:16]:
        for cell in (_token(c) for c in line.split("|")):
            if cell and any(cell == n or n in cell for n in names):
                return cell
    return ""


def _count_rows(preview: str) -> int:
    lines = [ln for ln in str(preview or "").split("\n") if "|" in ln and not re.match(r"^sheet:", ln.strip(), re.I)]
    return max(0, len(lines) - 1)


def _title(preview: str) -> str:
    match = re.search(r"job title\s*[:|#]\s*([^\n|]{3,80})", str(preview or ""), re.I)
    return re.sub(r"\s+", " ", match.group(1)).strip() if match else ""


def placement(kind: str, why: str, **extra) -> dict:
    return {"kind": kind, "category": KIND_TO_CATEGORY[kind], "placeIn": KIND_TO_PLACE[kind], "why": why,
            **{k: v for k, v in extra.items() if v}}


def place_file(file_name: str, preview: str = "") -> dict:
    name = _norm(file_name)
    collapsed = name.replace(" ", "_")
    preview = str(preview or "")
    blob = f"{name}\n{preview}".lower()
    emp_no = _find_column(preview, ["emp no", "employee id", "emp id", "employee code"])
    emp_name = _find_column(preview, ["emp name", "employee name", "full name"])
    skills = _find_column(preview, ["detailed skills", "skills bucket", "top 3 skills", "skills"])
    br_id = _find_column(preview, ["auto req id", "br id", "requisition id"])
    assigned_q = bool(re.search(r"assigned question\s*\d+", preview, re.I))
    jd_prose = bool(re.search(r"job title\s*:|mandatory skills\s*:|job description|primary skills\s*:", preview, re.I))
    if "user credential" in name or "employee user credential" in name:
        return placement("unknown", "Credential workbooks are not uploaded. Use Resource_Question_Mapping.xlsx only.")
    if assigned_q or "resource question mapping" in name or "question mapping" in name:
        return placement("portal_mapping", "This workbook maps people to portal test questions, not Corp Pool.")
    spreadsheet = bool(re.search(r"\.(xlsx|xls|csv)$", name))
    jd_doc = bool(re.search(r"\.(docx|doc|pdf|txt|html|htm)$", name))
    has_br = bool(re.search(r"\d+\s*br\b", name) or re.search(r"\d{5}br\b", name.replace(" ", "")))
    preview_br = bool(re.search(r"\b\d{5}\s*br\b", preview, re.I))
    if jd_doc and has_br:
        return placement("jd", "Filename has a BR ID on a JD document.", suggestedTitle=_title(preview))
    if spreadsheet and (has_br or "br_rawdata" in collapsed):
        return placement("br", "BR / requisition workbook (Auto req ID rows).", columns={"br_id": br_id} if br_id else None)
    if spreadsheet and (br_id or preview_br) and not emp_no:
        return placement("br", "Excel has an Auto req ID / BR ID column.", columns={"br_id": br_id} if br_id else None)
    if "corp pool" in name or "active list" in name or "employee list" in name or (emp_no and emp_name):
        return placement("corp_pool", "Roster columns: Emp No and name." if emp_no else "Filename matches a Corp Pool roster.",
                         peopleCount=_count_rows(preview),
                         columns={"emp_no": emp_no or "emp no", "name": emp_name or "emp name", "skills": skills or ""})
    if jd_prose and not emp_no:
        return placement("jd", "File reads as a job description (title / mandatory skills).", suggestedTitle=_title(preview))
    if name.endswith(".csv") and (emp_no or re.search(r"emp no|employee id", blob)):
        return placement("corp_pool", "CSV roster with employee IDs.", columns={"emp_no": emp_no or "emp no", "name": emp_name, "skills": skills})
    return placement("unknown", "Could not tell Corp Pool, JD/BR, or Portal Mapping from the filename or preview.")


def _cell_text(value) -> str:
    return "" if value is None else str(value).replace("\u00a0", " ").strip()


def excel_preview(data: bytes, max_rows: int = 12) -> str:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as err:
        log.warning("Excel preview unavailable: %s", err)
        return ""
    lines: list[str] = []
    for sheet in workbook.worksheets[:3]:
        lines.append(f"Sheet: {sheet.title}")
        count = 0
        for row in sheet.iter_rows(min_row=1, max_col=9, values_only=True):
            if count >= max_rows:
                break
            cells = [_cell_text(v) for v in row]
            if any(cells):
                lines.append(" | ".join(cells))
                count += 1
    return "\n".join(lines)[:7000]


def classify_with_llm(file_name: str, preview: str) -> dict:
    rules = place_file(file_name, preview)
    if rules["kind"] != "unknown" or not llm.is_up():
        return rules
    prompt = f"""You classify files for an HR screening console with two people lists that must stay separate:
- Corp Pool = bench/employees to match against job descriptions (Emp No, name, grade, skills)
- Employee Portal mapping = test/question mapping, NOT Corp Pool
- JD / BR = job descriptions / requisitions (role title, responsibilities, mandatory skills). A JD Excel is often one column of prose, not Emp No rows.

Filename: {file_name}

File contents (preview):
{preview[:6500]}

Return ONLY JSON:
{{"kind": "corp_pool" | "jd" | "br" | "portal_mapping" | "unknown", "why": "one or two sentences", "peopleCount": 0, "suggestedTitle": "", "columns": {{"emp_no": "", "name": "", "skills": ""}}}}
If it is a JD, suggestedTitle is the job title. columns only for corp_pool."""
    try:
        parsed = llm.complete_json(prompt, temperature=0.1, max_tokens=220, timeout_s=30)
        kind = parsed.get("kind") if parsed.get("kind") in KIND_TO_CATEGORY else "unknown"
        columns = parsed.get("columns") if isinstance(parsed.get("columns"), dict) else None
        count = parsed.get("peopleCount")
        return placement(kind, str(parsed.get("why") or "").strip() or "Could not classify this file.",
                         peopleCount=count if isinstance(count, int) else None,
                         suggestedTitle=str(parsed.get("suggestedTitle") or "").strip(), columns=columns)
    except Exception:
        return rules
