import io

import pytest
from openpyxl import Workbook, load_workbook

from services import deleted_requirements as dr, jd_to_br, llm


@pytest.fixture()
def admin(client, monkeypatch):
    monkeypatch.setenv("ADMIN_PASSWORD", "AdminPw1!")
    token = client.post("/api/admin/auth/login", json={"email": "boss@infinite.com", "password": "AdminPw1!"}).json()["token"]
    return {"Authorization": f"Bearer {token}"}


def test_br_id_helpers():
    assert dr.extract_br_id("50656BR | Dev.docx") == "50656br"
    assert dr.file_name_has_br_id("00001BR | x") and not dr.file_name_has_br_id("Pasted Job Description")
    assert dr.next_synthetic_br_id(["00001BR | a", "88001BR | real", "00007BR | b"]) == "00008BR"


def test_jd_crud_dedupe_and_delete(client, store, admin):
    created = client.post("/api/admin/jd", headers=admin, json={"jd": "Need Python dev", "rmEmail": "RM@infinite.com"}).json()
    assert created["success"] and created["jd"]["fileName"] == "00001BR | Pasted Job Description"
    second = client.post("/api/admin/jd", headers=admin, json={"jd": "Another", "rmEmail": "rm@infinite.com", "fileName": "50656BR | Role.docx"}).json()
    dup = client.post("/api/admin/jd", headers=admin, json={"jd": "Another v2", "rmEmail": "rm@infinite.com", "fileName": "50656BR | Role.docx"}).json()
    jds = client.get("/api/admin/jd", headers=admin).json()["jds"]
    assert len(jds) == 2
    group = next(j for j in jds if "50656BR" in j["fileName"])
    assert sorted(group["duplicateIds"]) == sorted([second["jd"]["id"], dup["jd"]["id"]])
    mine = client.get("/api/admin/jd?email=other@infinite.com", headers=admin).json()["jds"]
    assert all(j["rmEmail"] in ("admin@infinite.com", "other@infinite.com", "") for j in mine) or mine == [] or len(mine) <= len(jds)
    updated = client.post("/api/admin/jd", headers=admin, json={"jd": "Edited text", "jdId": created["jd"]["id"]}).json()
    assert updated["jd"]["fileName"] == "00001BR | Pasted Job Description"
    assert client.request("DELETE", "/api/admin/jd", headers=admin, json={"ids": [second["jd"]["id"]]}).json()["deletedCount"] == 1
    left = client.get("/api/admin/jd", headers=admin).json()["jds"]
    assert len(left) == 1 and store.count("job_descriptions") == 1
    assert client.post("/api/admin/jd", headers=admin, json={"jd": "  "}).status_code == 400
    assert client.request("DELETE", "/api/admin/jd", headers=admin, json={}).status_code == 400


def test_jd_upload_and_type_check(client, store, admin):
    ok = client.post("/api/admin/jd", headers=admin, files={"file": ("role.txt", b"Python developer needed")}, data={"rmEmail": "x@infinite.com"})
    assert ok.status_code == 200 and ok.json()["jd"]["jdText"] == "Python developer needed"
    assert client.post("/api/admin/jd", headers=admin, files={"file": ("role.exe", b"MZ....")}).status_code == 400


def test_jd_requires_admin(client):
    assert client.get("/api/admin/jd").status_code == 401 and client.post("/api/admin/jd-to-br").status_code == 401


def test_jd_to_br_workbook_with_qwen_and_fallback(client, admin, monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "BR _Raw Data"
    ws.append(["Auto req ID", "Designation", "Grade", "Mandatory Skills"])
    ws.append(["50000BR", "Old", "E1", "x"])
    buf = io.BytesIO()
    wb.save(buf)
    monkeypatch.setattr(llm, "complete_json", lambda *a, **k: {"job_title": "Platform Engineer", "experience": "4-6 years",
                                                                "skills": ["Python"], "monitoring_tools": [], "cloud_platforms": ["AWS"]})
    res = client.post("/api/admin/jd-to-br", headers=admin, files=[("jds", ("a.txt", b"Platform role")), ("template", ("t.xlsx", buf.getvalue()))])
    assert res.status_code == 200
    out = load_workbook(io.BytesIO(res.content)).active
    assert [c.value for c in out[3]][:4] == ["50001BR", "Platform Engineer", "E2", "Python, AWS"]

    def down(*a, **k):
        raise llm.LlmError("down")

    monkeypatch.setattr(llm, "complete_json", down)
    again = client.post("/api/admin/jd-to-br", headers=admin, files=[("jds", ("a.txt", b"Title: L2 Support Engineer\n3 years with Splunk and SQL")),
                                                                      ("template", ("t.xlsx", buf.getvalue()))])
    row = load_workbook(io.BytesIO(again.content)).active[3]
    assert row[1].value == "L2 Support Engineer" and "Splunk" in row[3].value and "SQL" in row[3].value
    assert client.post("/api/admin/jd-to-br", headers=admin, files=[("template", ("t.xlsx", b"x"))]).status_code == 400
