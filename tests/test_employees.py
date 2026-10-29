import io
import json

import pytest
from openpyxl import Workbook

from services import corp_pool, docs_storage, portal_roster as pr

EMP = "aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa"


@pytest.fixture()
def admin(client, monkeypatch, tmp_path):
    monkeypatch.setattr("services.storage.UPLOAD_DIR", tmp_path)
    monkeypatch.setattr("services.docs_storage.UPLOAD_DIR", tmp_path)
    monkeypatch.setattr("services.structured_log.UPLOAD_DIR", tmp_path)
    monkeypatch.setenv("ADMIN_PASSWORD", "AdminPw1!")
    pr.invalidate_mapping_cache()
    token = client.post("/api/admin/auth/login", json={"email": "boss@infinite.com", "password": "AdminPw1!"}).json()["token"]
    return {"Authorization": f"Bearer {token}"}


def roster(*ids):
    return [{"employee_id": i, "full_name": f"Name {i}", "skills": "python", "department": "eng", "designation": "dev", "status": "Active",
             "grade": "E1", "score": 10, "shortlisted": False, "matchingSkills": []} for i in ids]


def mapping_book() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Emp ID", "Emp Name", "Product", "Assigned Question 1", "Assigned Question 2", "Assigned Question 3"])
    ws.append(["1001", "Mapped One", "One-NDS", "What is X?", "What is Y?", "what is x?"])
    ws.append(["9999", "Mapping Only", "Prod", "Q", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_requires_admin(client):
    assert client.get("/api/admin/employees").status_code == 401
    assert client.get("/api/admin/employees/mapping-questions?employeeId=1").status_code == 401


def test_roster_crud_and_shortlist(client, store, admin):
    corp_pool.save_roster(roster("E1", "E2", "E3"))
    data = client.get("/api/admin/employees", headers=admin).json()
    assert [e["employee_id"] for e in data["employees"]] == ["E1", "E2", "E3"]
    one = client.post("/api/admin/employees", headers=admin, json={"employeeId": "E1"}).json()
    assert one["employee"]["shortlisted"] is True
    bulk = client.post("/api/admin/employees", headers=admin, json={"employeeIds": ["E1", "E2"], "shortlisted": True}).json()
    assert bulk["updatedCount"] == 2 and all(e["shortlisted"] for e in bulk["employees"])
    mixed = client.post("/api/admin/employees", headers=admin, json={"changes": [{"employeeId": "E2", "shortlisted": False}]}).json()
    assert mixed["employee"]["shortlisted"] is False
    assert client.post("/api/admin/employees", headers=admin, json={"employeeId": "ZZ"}).status_code == 404
    patched = client.patch("/api/admin/employees", headers=admin,
                           json={"employeeId": "E3", "updates": {"full_name": "New Name", "employee_id": "E9"}}).json()
    assert patched["employee"]["employee_id"] == "E9" and patched["employee"]["manually_edited"] is True
    assert client.patch("/api/admin/employees", headers=admin,
                        json={"employeeId": "E9", "updates": {"score": 150}}).status_code == 400
    assert client.patch("/api/admin/employees", headers=admin, json={"employeeId": "E1", "updates": {"employee_id": "E2"}}).status_code == 409
    gone = client.request("DELETE", "/api/admin/employees", headers=admin, json={"ids": ["E1"]}).json()
    assert gone["deleted"] == 1 and [e["employee_id"] for e in gone["employees"]] == ["E2", "E9"]
    assert "e1" in corp_pool.load_deleted()["ids"]
    assert client.request("DELETE", "/api/admin/employees", headers=admin, json={}).status_code == 400
    assert client.request("DELETE", "/api/admin/employees", headers=admin, json={"ids": ["NOPE"]}).status_code == 404


def test_csv_export(client, store, admin):
    corp_pool.save_roster(roster("E1"))
    res = client.get("/api/admin/employees?export=true", headers=admin)
    assert res.headers["content-type"].startswith("text/csv") and res.text.splitlines()[0].startswith("Employee ID,Name")
    assert "E1,Name E1,eng,dev,python,Active,E1,10,No" in res.text


def test_portal_employees_merge_mapping_and_tests(client, store, admin):
    docs_storage.write_file("Portal Mapping", "Resource_Question_Mapping.xlsx", mapping_book())
    store.seed("employees", [{"id": EMP, "employee_id": "1001", "email": "a@x.com", "full_name": "Acct One", "role": "employee",
                              "department": "eng", "product": "One-NDS", "product_qb_eligible": True}])
    store.seed("tests", [{"id": "t1", "employee_id": EMP, "employee_code": "1001", "topic_id": "resource-product-assessment", "subject_id": "s",
                          "topic_title": "Prod", "total_questions": 25, "status": "completed", "score_correct": 20, "score_percent": 80,
                          "completed_at": "2026-09-30T10:00:00Z", "proctoring": json.dumps({"warningCount": 1})}])
    data = client.get("/api/admin/employees", headers=admin).json()
    portal = {r["employee_id"]: r for r in data["resourcePortalEmployees"]}
    one = portal["1001"]
    assert one["product"] == "One-SDL" and one["assigned_questions"] == ["What is X?", "What is Y?"] and one["assigned_question_count"] == 2
    assert one["test_status"] == "completed" and one["score"] == 20 and one["test_id"] == "t1"
    assert portal["9999"]["test_status"] == "not_started" or portal["9999"]["test_status"] == "unassigned"
    assert data["allTestResults"][0]["scorePercent"] == 80 and data["allTestResults"][0]["hasRecording"] is False
    mapped = client.get("/api/admin/employees/mapping-questions?employeeId=1001.0", headers=admin).json()
    assert mapped["assigned_questions"] == ["What is X?", "What is Y?"]
    assert client.get("/api/admin/employees/mapping-questions", headers=admin).status_code == 400


def test_portal_hidden_for_screening_only_admin(client, store, admin):
    corp_pool.save_roster(roster("E1"))
    store.seed("tests", [{"id": "t1", "employee_id": EMP, "employee_code": "1001", "topic_id": "x", "subject_id": "s", "status": "pending", "total_questions": 5}])
    hidden = client.get("/api/admin/employees?email=ramendras@infinite.com", headers=admin).json()
    assert hidden["allTestResults"] == [] and hidden["resourcePortalEmployees"] == [] and len(hidden["employees"]) == 1
    visible = client.get("/api/admin/employees?email=boss@infinite.com", headers=admin).json()
    assert len(visible["allTestResults"]) == 1


def test_status_derivation():
    assert pr.derive_status(0, None, None, 0, 0, None) == "unassigned"
    assert pr.derive_status(25, "t", "pending", 5, 25, None) == "pending"
    assert pr.derive_status(25, "t", "in_progress", 0, 25, "2026") == "in_progress"
    assert pr.derive_status(25, "t", "completed", 25, 25, None) == "completed"
    assert pr.normalize_employee_id("1001.0") == "1001" and pr.mapping_looks_valid(mapping_book()) is True


def test_portal_roster_identifies_password_managed_accounts(store):
    store.seed("employees", [
        {"id": "e1", "employee_id": "1001", "email": "a@x.com", "full_name": "Managed",
         "role": "employee", "product": "One-NDS", "product_qb_eligible": True},
    ])
    eligible = pr._eligible_accounts()
    assert len(eligible) == 1 and eligible[0]["product_qb_eligible"] is True

    mapping = {
        "employee_id": "1001", "full_name": "Mapped", "role": "employee", "domain": "eng",
        "product": "One-NDS", "email": "a@x.com", "ddh": "", "emp_status": "Confirmed",
        "remarks": "", "assigned_questions": [], "assigned_question_count": 0,
    }
    assert pr._merge(eligible[0], mapping)["product_qb_eligible"] is True
    assert pr._merge({**mapping, "product_qb_eligible": False}, mapping)["product_qb_eligible"] is False


def test_reset_test_and_delete_video(client, store, admin):
    from services import test_video

    tid, q = "bbbbbbbb-bbbb-4bbb-8bbb-bbbbbbbbbbbb", "cccccccc-cccc-4ccc-8ccc-cccccccccc00"
    store.seed("employees", [{"id": EMP, "employee_id": "1001", "email": "a@x.com", "full_name": "A", "xp_points": 50, "ai_readiness_score": 60}])
    store.seed("tests", [{"id": tid, "employee_id": EMP, "employee_code": "1001", "topic_id": pr.PRODUCT_TOPIC, "subject_id": "s", "topic_title": "Prod",
                          "status": "completed", "score_percent": 80, "total_questions": 1, "completed_at": "2026-09-01T00:00:00Z"}])
    store.seed("test_questions", [{"id": q, "test_id": tid, "question_index": 0, "question_text": "Q", "options": ["a", "b"], "correct_option_index": 0}])
    store.seed("test_attempts", [{"id": "a1", "test_id": tid, "employee_id": EMP, "question_id": q, "is_correct": True}])
    webm = b"\x1a\x45\xdf\xa3" + b"\x1f\x43\xb6\x75" + b"v" * 6000
    assert test_video.save(tid, webm)
    gone = client.post("/api/admin/employees/delete-video", headers=admin, json={"employeeId": "1001"}).json()
    assert gone["success"] and not test_video.exists(tid) and store.first("tests", {"id": tid})["status"] == "completed"
    bulk = client.post("/api/admin/employees/delete-video", headers=admin, json={"items": [{"testId": tid}, {"employeeId": "ghost"}]}).json()
    assert bulk["deletedCount"] == 1 and bulk["skippedCount"] == 1
    assert client.post("/api/admin/employees/delete-video", headers=admin, json={}).status_code == 400
    res = client.post("/api/admin/employees/reset-test", headers=admin, json={"employeeId": "1001"}).json()
    assert res == {"success": True, "testId": tid, "preservedQuestions": 1}
    assert store.first("tests", {"id": tid})["status"] == "pending" and store.count("test_attempts", {"test_id": tid}) == 0
    assert store.count("tests") == 2
    assert store.first("employees", {"employee_id": "1001"})["xp_points"] == 0
    assert client.post("/api/admin/employees/reset-test", headers=admin, json={"employeeId": "ghost"}).status_code == 404
    assert client.post("/api/admin/employees/reset-test", headers=admin, json={}).status_code == 400


def test_dispatch_mail_corp_pool_and_portal(client, store, admin):
    rows = roster("E1", "E2")
    rows[0].update(shortlisted=True, email="e1@x.com", full_name="<b>Eve</b>")
    rows[1].update(email="e2@x.com")
    corp_pool.save_roster(rows)
    res = client.post("/api/admin/employees/dispatch_mail", headers=admin, json={"adminEmail": "rm@infinite.com"}).json()
    assert res["count"] == 1 and res["dispatched"][0]["status"] == "simulated"
    mail = store.select("simulated_emails")[0]
    assert "&lt;b&gt;Eve&lt;/b&gt;" in mail["body"] and "<b>Eve</b>" not in mail["body"] and mail["rm_email"] == "rm@infinite.com"
    assert store.first("candidate_sessions", {"email": "e1@x.com"}) is not None and store.count("resumes") == 1
    portal = client.post("/api/admin/employees/dispatch_mail", headers=admin,
                         json={"portal": True, "recipients": [{"employee_id": "1001", "email": "p@x.com", "full_name": "Pat"}, {"employee_id": "", "email": "x@x.com"}]}).json()
    assert portal["count"] == 1 and store.count("resumes") == 1
    assert "/employee?org=BizX" in store.select("simulated_emails")[-1]["body"] and "Employee ID: <strong>1001</strong>" in store.select("simulated_emails")[-1]["body"]
    assert client.post("/api/admin/employees/dispatch_mail", headers=admin, json={"employeeIds": ["E2"]}).status_code == 400
    corp_pool.save_roster([])
    assert client.post("/api/admin/employees/dispatch_mail", headers=admin, json={}).status_code == 404


def test_jd_match_via_shared_engine(client, store, admin):
    import shutil

    if not shutil.which("node"):
        pytest.skip("node not installed")
    jd_id = "dddddddd-dddd-4ddd-8ddd-dddddddddddd"
    store.seed("job_descriptions", [{"id": jd_id, "file_name": "1BR | Java", "jd_text": "Job Title: Java Developer\nMandatory Skills: Java, Spring Boot, SQL, Docker\n3-6 years"}])
    rows = roster("E1", "E2")
    rows[0].update(skills="Java, Spring Boot, MySQL, Docker, Git", designation="Senior Software Engineer", grade="E2")
    rows[1].update(skills="Python, React", designation="QA Engineer", grade="E1")
    corp_pool.save_roster(rows)
    listed = client.get(f"/api/admin/employees?activeJdId={jd_id}", headers=admin).json()["employees"]
    assert listed[0]["score"] > listed[1]["score"] and "Spring" in listed[0]["matchingSkills"]
    csv_text = client.get(f"/api/admin/employees?export=true&activeJdId={jd_id}", headers=admin).text
    assert f",{listed[0]['score']},No" in csv_text
    res = client.post("/api/admin/employees/llm-evaluate", headers=admin, json={"jdId": jd_id, "employeeIds": ["E1", "E2"], "scanId": "s1"}).json()
    assert res["scored"] == 2 and res["done"] is True and res["results"][0]["engine"] == "chips"
    saved = {e["employee_id"]: e for e in corp_pool.load_roster()}
    assert saved["E1"]["score_override_jd_id"] == jd_id and saved["E1"]["score"] == listed[0]["score"]
    again = client.post("/api/admin/employees/llm-evaluate", headers=admin, json={"jdId": jd_id, "employeeIds": ["E1"], "force": False, "scanId": "s2"}).json()
    assert again["skipped"] == 1 and again["scored"] == 0
    override = client.get(f"/api/admin/employees?activeJdId={jd_id}", headers=admin).json()["employees"][0]
    assert override["score"] == saved["E1"]["score_override"]
    assert client.post("/api/admin/employees/llm-evaluate", headers=admin, json={"jdId": "all", "employeeIds": ["E1"]}).status_code == 400
    assert client.post("/api/admin/employees/llm-evaluate", headers=admin, json={"jdId": jd_id}).status_code == 400


def test_exports(client, store, admin):
    from openpyxl import load_workbook

    store.seed("employees", [{"id": EMP, "employee_id": "1001", "email": "a@x.com", "full_name": "Acct One", "role": "employee", "product": "One-NDS",
                              "product_qb_eligible": True}])
    tid, q = "bbbbbbbb-bbbb-4bbb-8bbb-bbbbbbbbbbbb", "cccccccc-cccc-4ccc-8ccc-cccccccccc00"
    store.seed("tests", [{"id": tid, "employee_id": EMP, "employee_code": "1001", "topic_id": pr.PRODUCT_TOPIC, "subject_id": "s", "topic_title": "Prod",
                          "total_questions": 2, "status": "completed", "score_correct": 1, "score_percent": 50, "completed_at": "2026-09-30T10:00:00Z",
                          "proctoring": {"warningCount": 1, "violations": [{"type": "Tab Switch", "timestamp": "2026-09-30T09:00:00Z"}]}}])
    store.seed("test_questions", [{"id": q, "test_id": tid, "question_index": 0, "question_text": "Q one?", "options": ["a", "b"], "correct_option_index": 1}])
    store.seed("test_attempts", [{"id": "a1", "test_id": tid, "employee_id": EMP, "question_id": q, "selected_option_index": 1, "is_correct": True}])
    res = client.get("/api/admin/employees/export-portal", headers=admin)
    assert res.status_code == 200 and "employee_portal_test_results_" in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    top = [c.value for c in wb["Employee Portal Results"][2]]
    assert top[0] == "Acct One" and top[4] == "One-SDL" and top[7] == "Completed" and top[9] == "1/2 (50%)" and top[10] == "1" and top[11].startswith("1. Tab Switch")
    detail = [c.value for c in wb["Question Details"][2]]
    assert detail[:6] == ["1001", "Acct One", 1, "Q one?", "B) b", "Correct"]
    assert client.get("/api/admin/employees/export-portal?email=ramendras@infinite.com", headers=admin).status_code == 403

    short = client.post("/api/admin/employees/export-shortlisted", headers=admin,
                        json={"employees": [{"employee_id": "E1", "full_name": "Eve", "score": 71.6, "matchDecision": "interview", "matchingSkills": ["Java", "SQL"]}],
                              "fileName": "my:list"})
    assert 'filename="my-list.xlsx"' in short.headers["content-disposition"]
    row = [c.value for c in load_workbook(io.BytesIO(short.content)).active[2]]
    assert row[0] == "Eve" and row[6] == "72%" and row[7] == "INTERVIEW" and row[8] == "Java, SQL" and row[10] == "Shortlisted"
    assert client.post("/api/admin/employees/export-shortlisted", headers=admin, json={"employees": []}).status_code == 400

    book = io.BytesIO()
    from openpyxl import Workbook

    w = Workbook()
    w.active.append(["Emp No", "Match Score"])
    for s in (10, 40, 60, 80, "90%"):
        w.active.append(["x", s])
    w.save(book)
    pct = client.post("/api/admin/employees/percentage-matching", headers=admin,
                      files=[("files", ("51210BR vs pool.xlsx", book.getvalue())), ("files", ("b.csv", b"name,score\nA,20\nB,95\n"))])
    ws = load_workbook(io.BytesIO(pct.content)).active
    assert [c.value for c in ws[2]][2:] == [1, 1, 1, 2] and [c.value for c in ws[3]][2:] == [1, 0, 0, 1]
    assert client.post("/api/admin/employees/percentage-matching", headers=admin, files=[("files", ("c.csv", b"name\nA\n"))]).status_code == 400
    assert client.post("/api/admin/employees/percentage-matching", headers=admin).status_code == 400
