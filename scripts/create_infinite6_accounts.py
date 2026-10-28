"""Create Infinite tester accounts + assign CFX product assessment on Live."""
from __future__ import annotations

import os
import random
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ.setdefault("MSSQL_TARGET", "remote")
os.environ.setdefault("MSSQL_REMOTE_DATABASE", "TalentScope_DB_Live")

HANDOFF_SCRIPTS = (
    Path(__file__).resolve().parents[2] / "Handoff" / "test schedule file" / "scripts"
)
sys.path.insert(0, str(HANDOFF_SCRIPTS))

from openpyxl import load_workbook  # noqa: E402

from assign_product_tests import assign_assessment  # noqa: E402
from core.config import mssql_settings  # noqa: E402
from core.db import MssqlStore, set_store  # noqa: E402
from services import accounts  # noqa: E402

QB = (
    Path(__file__).resolve().parents[2]
    / "Handoff"
    / "test schedule file"
    / "question-banks"
    / "Question_Bank_Filled.xlsx"
)
PRODUCT = "CFX"
PER_TEST = 25
TOPIC_ID = accounts.PRODUCT_ASSESSMENT_TOPIC_ID

# Emp ID #5 (ravitejake) was not provided in the request — skipped.
PEOPLE = [
    {"employee_id": "1031886", "email": "chandankum@infinite.com", "full_name": "Chandan Kumar"},
    {"employee_id": "1022211", "email": "nikhilh@infinite.com", "full_name": "Nikhil H"},
    {"employee_id": "1029830", "email": "poojash@infinite.com", "full_name": "Pooja S H"},
    {"employee_id": "1029930", "email": "sapriyanka@infinite.com", "full_name": "Priyanka S A"},
    {"employee_id": "1032657", "email": "anshikas@infinite.com", "full_name": "Anshika S"},
]


def load_cfx_wide(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb["CFX"]
    rows = list(ws.iter_rows(values_only=True))
    letter = {"A": 0, "B": 1, "C": 2, "D": 3}
    out = []
    for raw in rows[1:]:
        if not raw or raw[3] is None:
            continue
        text = str(raw[3]).strip()
        options = [str(raw[i] or "").strip() for i in range(4, 8)]
        ans = str(raw[8] or "").strip().upper()[:1]
        if ans not in letter or not text or not all(options):
            continue
        out.append(
            {
                "question_text": text,
                "options": options,
                "correct_option_index": letter[ans],
                "explanation": "CFX product assessment.",
                "difficulty": "medium",
            }
        )
    return out


def ensure_account(row: dict) -> dict:
    existing = accounts.get_account(row["employee_id"]) or accounts.get_account_by_email(row["email"])
    pwd_hash, pwd_salt = accounts.hash_password(row["employee_id"])
    account = {
        "employee_id": row["employee_id"],
        "full_name": row["full_name"],
        "email": row["email"].strip().lower(),
        "department": (existing or {}).get("department") or "general",
        "role": "employee",
        "product": PRODUCT,
        "product_qb_eligible": True,
        "assessment_only": True,
        "is_first_login": True,
        "password_hash": pwd_hash,
        "password_salt": pwd_salt,
        "xp_points": (existing or {}).get("xp_points", 0),
        "streak_days": (existing or {}).get("streak_days", 0),
        "skill_level": (existing or {}).get("skill_level") or "beginner",
        "ai_readiness_score": (existing or {}).get("ai_readiness_score", 0),
    }
    if not accounts.upsert_account(account):
        raise RuntimeError(f"upsert failed {row['employee_id']}")
    saved = accounts.get_account(row["employee_id"])
    if not saved:
        raise RuntimeError(f"missing after upsert {row['employee_id']}")
    return saved


def verify_person(store: MssqlStore, row: dict) -> dict:
    by_id = accounts.get_account(row["employee_id"])
    by_email = accounts.get_account_by_email(row["email"])
    emp_uuid = accounts.employee_uuid(row["employee_id"]) if by_id else None
    test = None
    qcount = 0
    if emp_uuid:
        test = store.first(
            "tests",
            {"employee_id": emp_uuid, "topic_id": TOPIC_ID},
            ["id", "topic_title", "status", "total_questions"],
        )
        if test:
            qcount = store.query(
                "SELECT COUNT(*) AS c FROM dbo.test_questions WHERE test_id = ?",
                (test["id"],),
            )[0]["c"]
    pw_ok = bool(
        by_id
        and accounts.verify_password(
            row["employee_id"], by_id.get("password_salt") or "", by_id.get("password_hash") or ""
        )
    )
    return {
        "account_ok": bool(by_id),
        "email_ok": bool(by_email) and by_email["employee_id"] == row["employee_id"],
        "sso_email": by_id["email"] if by_id else "",
        "pw_emp_id": pw_ok,
        "first_login": bool(by_id and by_id.get("is_first_login")),
        "test_topic": (test or {}).get("topic_title") or "",
        "test_status": (test or {}).get("status") or "",
        "qcount": int(qcount or 0),
    }


def main() -> int:
    set_store(None)
    store = MssqlStore(mssql_settings())
    set_store(store)
    print("DB", mssql_settings()["database"], "ping", store.ping(), flush=True)

    bank = load_cfx_wide(QB)
    print("CFX bank", len(bank), flush=True)
    if len(bank) < 2:
        print("ERROR: empty CFX bank")
        return 2

    for row in PEOPLE:
        account = ensure_account(row)
        rng = random.Random(f"{row['employee_id']}:CFX:20261027")
        picked = bank[:] if len(bank) <= PER_TEST else rng.sample(bank, PER_TEST)
        result = assign_assessment(store, account, picked, PRODUCT)
        print(
            f"ASSIGNED {row['employee_id']} {row['email']} q={len(picked)} {result}",
            flush=True,
        )

    print("\n=== VERIFY ===", flush=True)
    all_ok = True
    for row in PEOPLE:
        v = verify_person(store, row)
        ok = v["account_ok"] and v["email_ok"] and v["pw_emp_id"] and v["qcount"] > 0
        all_ok = all_ok and ok
        print(
            f"{'OK' if ok else 'FAIL'}\t{row['employee_id']}\t{row['email']}\t"
            f"acct={v['account_ok']}\temail={v['email_ok']}\tpw={v['pw_emp_id']}\t"
            f"sso={v['sso_email']}\ttopic={v['test_topic']}\tq={v['qcount']}\tstatus={v['test_status']}",
            flush=True,
        )
    print("SKIPPED ravitejake@infinite.com — Emp ID #5 not provided in request", flush=True)
    return 0 if all_ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
