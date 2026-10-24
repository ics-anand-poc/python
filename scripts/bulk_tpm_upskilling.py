"""Bulk-create TPM Upskilling portal accounts + assign MCQ product assessment.

Reads:
  - TPM credentials Excel (Emp ID, Name, Email, Role)
  - TPM MCQs Excel (Question, Option A-D, Correct Answer A-D)

Usage (from python/):
  set MSSQL_TARGET=remote
  set MSSQL_REMOTE_DATABASE=TalentScope_DB_Dev
  venv\\Scripts\\python.exe scripts\\bulk_tpm_upskilling.py
"""
from __future__ import annotations

import argparse
import os
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from core.config import mssql_settings
from core.db import MssqlStore, set_store
from services import accounts

PRODUCT = "TPM Upskilling"
TOPIC_ID = accounts.PRODUCT_ASSESSMENT_TOPIC_ID
SUBJECT_ID = "resource-subject"
SUBJECT_TITLE = "Product Assessment"
TIME_LIMIT_SECONDS = 2700  # 45 min for 35 questions


def load_credentials(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for raw in rows[1:]:
        if not raw or raw[0] is None:
            continue
        emp_id = str(raw[0]).strip()
        name = str(raw[1] or emp_id).strip()
        email = str(raw[2] or "").strip().lower()
        role = str(raw[3] or "TPM").strip()
        if not emp_id or not email:
            continue
        out.append({"employee_id": emp_id, "full_name": name, "email": email, "role_label": role})
    return out


def load_mcqs(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    letter_to_idx = {"A": 0, "B": 1, "C": 2, "D": 3}
    out = []
    for raw in rows[1:]:
        if not raw or raw[1] is None:
            continue
        q_no = raw[0]
        text = str(raw[1]).strip()
        options = [str(raw[i] or "").strip() for i in range(2, 6)]
        correct_letter = str(raw[6] or "").strip().upper()[:1]
        if correct_letter not in letter_to_idx:
            raise ValueError(f"Bad correct answer for Q{q_no}: {raw[6]!r}")
        if not text or not all(options):
            raise ValueError(f"Incomplete MCQ row Q{q_no}")
        out.append(
            {
                "question_text": text,
                "options": options,
                "correct_option_index": letter_to_idx[correct_letter],
                "explanation": str(raw[7] or "").strip(),
                "difficulty": "medium",
            }
        )
    return out


def ensure_account(row: dict) -> dict:
    existing = accounts.get_account(row["employee_id"])
    pwd_hash = existing.get("password_hash") if existing else None
    pwd_salt = existing.get("password_salt") if existing else None
    if not pwd_hash or not pwd_salt:
        # Initial password = Employee ID (same convention as first-login Emp ID)
        pwd_hash, pwd_salt = accounts.hash_password(row["employee_id"])
        first_login = True
    else:
        first_login = bool(existing.get("is_first_login")) if existing else False

    account = {
        "employee_id": row["employee_id"],
        "full_name": row["full_name"],
        "email": row["email"],
        "department": (existing or {}).get("department") or "general",
        "role": "employee",
        "product": PRODUCT,
        "product_qb_eligible": True,
        "assessment_only": True,
        "is_first_login": first_login,
        "password_hash": pwd_hash,
        "password_salt": pwd_salt,
        "xp_points": (existing or {}).get("xp_points", 0),
        "streak_days": (existing or {}).get("streak_days", 0),
        "skill_level": (existing or {}).get("skill_level") or "beginner",
        "ai_readiness_score": (existing or {}).get("ai_readiness_score", 0),
    }
    if not accounts.upsert_account(account):
        raise RuntimeError(f"Failed upsert {row['employee_id']}")
    saved = accounts.get_account(row["employee_id"])
    if not saved:
        raise RuntimeError(f"Missing after upsert {row['employee_id']}")
    return saved


def assign_assessment(store: MssqlStore, account: dict, mcqs: list[dict]) -> str:
    emp_uuid = accounts.employee_uuid(account["employee_id"])
    if not emp_uuid:
        raise RuntimeError(f"No UUID for {account['employee_id']}")

    existing = store.first(
        "tests",
        {"employee_id": emp_uuid, "topic_id": TOPIC_ID},
        ["id", "topic_title", "status", "total_questions"],
    )
    now = datetime.now(timezone.utc).isoformat()
    if existing:
        test_id = existing["id"]
        store.execute("DELETE FROM dbo.test_attempts WHERE test_id = ?", (test_id,))
        store.execute("DELETE FROM dbo.test_questions WHERE test_id = ?", (test_id,))
        store.update(
            "tests",
            {
                "employee_code": account["employee_id"],
                "topic_title": PRODUCT,
                "subject_id": SUBJECT_ID,
                "subject_title": SUBJECT_TITLE,
                "difficulty": "medium",
                "total_questions": len(mcqs),
                "time_limit_seconds": TIME_LIMIT_SECONDS,
                "status": "pending",
                "current_question_index": 0,
                "started_at": None,
                "completed_at": None,
                "in_progress": None,
                "session_recording_url": None,
                "proctoring": None,
                "score_correct": None,
                "score_total": None,
                "score_percent": None,
                "ai_analysis": None,
            },
            {"id": test_id},
        )
        action = "replaced"
    else:
        test_id = str(uuid.uuid4())
        store.insert(
            "tests",
            {
                "id": test_id,
                "employee_id": emp_uuid,
                "employee_code": account["employee_id"],
                "topic_id": TOPIC_ID,
                "subject_id": SUBJECT_ID,
                "topic_title": PRODUCT,
                "subject_title": SUBJECT_TITLE,
                "difficulty": "medium",
                "total_questions": len(mcqs),
                "time_limit_seconds": TIME_LIMIT_SECONDS,
                "status": "pending",
                "current_question_index": 0,
                "created_at": now,
            },
        )
        action = "created"

    for idx, q in enumerate(mcqs):
        store.insert(
            "test_questions",
            {
                "id": str(uuid.uuid4()),
                "test_id": test_id,
                "question_index": idx,
                "question_text": q["question_text"],
                "options": q["options"],
                "correct_option_index": q["correct_option_index"],
                "explanation": q.get("explanation") or "TPM Upskilling MCQ (23 Oct 2026).",
                "difficulty": q.get("difficulty") or "medium",
                "topic_id": TOPIC_ID,
                "topic_title": PRODUCT,
                "created_at": now,
            },
        )
    return f"{action}:{test_id}"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--credentials",
        default=r"C:\Users\aryanmi\Downloads\TPM Crediantials to be created.xlsx",
    )
    parser.add_argument(
        "--mcqs",
        default=r"C:\Users\aryanmi\Downloads\TPM_Upskilling_MCQs- 23rd Oct 26.xlsx",
    )
    parser.add_argument("--database", default=os.environ.get("MSSQL_REMOTE_DATABASE") or "TalentScope_DB_Dev")
    args = parser.parse_args()

    os.environ["MSSQL_TARGET"] = "remote"
    os.environ["MSSQL_REMOTE_DATABASE"] = args.database
    set_store(None)
    settings = mssql_settings()
    store = MssqlStore(settings)
    set_store(store)
    if not store.ping():
        print(f"ERROR: cannot reach {settings['server']}/{settings['database']}")
        return 2

    creds = load_credentials(Path(args.credentials))
    mcqs = load_mcqs(Path(args.mcqs))
    print(f"DB={settings['database']}  people={len(creds)}  mcqs={len(mcqs)}  product={PRODUCT!r}")

    ok = fail = 0
    report = []
    for row in creds:
        try:
            account = ensure_account(row)
            result = assign_assessment(store, account, mcqs)
            ok += 1
            report.append(f"OK {row['employee_id']} {row['email']} {result}")
            print(report[-1])
        except Exception as err:
            fail += 1
            msg = f"FAIL {row['employee_id']} {row['email']}: {err}"
            report.append(msg)
            print(msg)

    out = Path(__file__).resolve().parents[2] / "TalentScope-SharePoint-Handoff" / f"TPM_Upskilling_report_{args.database}.txt"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(
        "\n".join(
            [
                f"TPM Upskilling bulk load — {args.database}",
                f"When: {datetime.now(timezone.utc).isoformat()}",
                f"Product: {PRODUCT}",
                f"MCQs: {len(mcqs)}",
                f"People OK/FAIL: {ok}/{fail}",
                "Initial password for NEW accounts (no prior password): Employee ID",
                "Existing passwords kept.",
                "",
                *report,
            ]
        ),
        encoding="utf-8",
    )
    print(f"Wrote {out}")
    print(f"DONE ok={ok} fail={fail}")
    return 0 if fail == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
